"""
Assertion Testing Tab - For search transactions and assertion retrieval
"""
import json
import os
import re
import subprocess
import sys
import tempfile
import tkinter as tk
from tkinter import ttk
from tkinter import filedialog, messagebox, simpledialog
import customtkinter as ctk
from pathlib import Path
from utils.excel_handler import ExcelHandler, validate_assertion_excel, create_excel_template
from utils.api_client import APIClient
from utils.variable_processor import VariableProcessor
import threading
from datetime import datetime
import shutil
from .theme import FONTS
from utils.azure_reporting import (
    create_ado_bug, scan_assertion_failures_by_tag,
    AVAILABLE_REPORT_TAGS, SYSTEM1_AVAILABLE_REPORT_TAGS,
    TARGET_TAGS, SYSTEM1_TARGET_TAGS,
    test_ado_connection, get_pbi_details, check_existing_bug, add_failure_comment,
    scan_all_results, scan_all_assertion_results, update_bug_for_qa_result,
    get_all_ready_for_qa_bugs, get_all_closed_bugs, match_bug_to_test_case,
)
from utils.csv_custom_report import generate_custom_combined_report, get_available_rules_summary_columns
from utils.env_vars import EnvVar
from utils.execution_mode import ExecutionMode


class ToolTip:
    def __init__(self, widget, text):
        self.widget = widget
        self.text = text
        self.tooltip_window = None
        self.widget.bind("<Enter>", self.show_tooltip)
        self.widget.bind("<Leave>", self.hide_tooltip)

    def show_tooltip(self, event=None):
        if self.tooltip_window:
            return
        x = self.widget.winfo_rootx() + 20
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 10
        self.tooltip_window = tk.Toplevel(self.widget)
        self.tooltip_window.wm_overrideredirect(True)
        frame = tk.Frame(self.tooltip_window, background="#ffffe0", relief="solid", borderwidth=1)
        frame.pack()
        tk.Label(frame, text=self.text, background="#ffffe0", justify="left",
                 font=("Segoe UI", 9)).pack(padx=6, pady=4)
        self.tooltip_window.wm_geometry(f"+{x}+{y}")

    def hide_tooltip(self, event=None):
        if self.tooltip_window:
            self.tooltip_window.destroy()
            self.tooltip_window = None


class AssertionTab:
    def __init__(self, parent):
        self.frame = ctk.CTkFrame(parent, fg_color="transparent")
        self.frame.pack(fill="both", expand=True)

        self.excel_handler = ExcelHandler()
        self.api_client = APIClient()
        self.variable_processor = VariableProcessor()
        self.testing_active = False
        self.sheet_names = []
        self.current_sheet = None
        self.current_process = None

        # Pagination variables
        self.all_data = [] 
        self.current_page = 0
        self.rows_per_page = 100
        self.total_pages = 0

        # Sheet data cache for all sheets
        self.all_sheets_data = {} 

        self.use_specific_login = None
        self.provider_value = None
        self.payer_value = None
        self.tpa_value = None
        self.pharmacy_value = None

        self.custom_report_sheets = []
        self.generate_custom_report = False
        self.last_custom_report_path = None
        self.last_allure_dir = None
        self.allure_report_generated = False

        self.create_widgets()

    def create_widgets(self):
        """Create assertion testing widgets"""
        # Create scrollable container
        self.main_frame = ctk.CTkScrollableFrame(self.frame)
        self.main_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # Description
        description_text = (
            "This tab is used to test all search transactions which requires posting some request, "
            "and asserting the retrieval of the request through the search API's"
        )
        ctk.CTkLabel(self.main_frame, text=description_text, wraplength=800, justify="center", font=FONTS["main"]).pack(anchor="center", pady=(0, 15))

        # File selection
        file_frame = ctk.CTkFrame(self.main_frame)
        file_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(file_frame, text="Excel File:", anchor="w", font=FONTS["main_bold"]).pack(side="left", padx=10)
        self.file_path_var = tk.StringVar()
        ctk.CTkEntry(file_frame, textvariable=self.file_path_var, width=400, font=FONTS["main"]).pack(side="left", fill="x", expand=True, padx=5)
        ctk.CTkButton(file_frame, text="Browse", command=self.browse_file, width=80, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left", padx=(0, 5))
        ctk.CTkButton(file_frame, text="Load", command=self.load_file, width=80, font=FONTS["button"]).pack(side="left", padx=5)

        # Sheet selection frame (initially hidden)
        self.sheet_frame = ctk.CTkFrame(self.main_frame)
        
        ctk.CTkLabel(self.sheet_frame, text="Select Sheet:", anchor="w", font=FONTS["main_bold"]).pack(side="left", padx=10)
        self.sheet_combo = ctk.CTkOptionMenu(self.sheet_frame, values=[], command=self.on_sheet_selected_optionmenu, width=200, font=FONTS["main"], dropdown_font=FONTS["main"])
        self.sheet_combo.pack(side="left", padx=5)
        ctk.CTkButton(self.sheet_frame, text="Refresh", command=self.refresh_sheet, width=80, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left", padx=5)

        # Pagination controls frame
        self.pagination_frame = ctk.CTkFrame(self.main_frame)
        
        left_controls = ctk.CTkFrame(self.pagination_frame, fg_color="transparent")
        left_controls.pack(side="left", fill="x", expand=True)

        ctk.CTkLabel(left_controls, text="Rows per page:", anchor="w", font=FONTS["main"]).pack(side="left", padx=10)
        self.rows_per_page_var = tk.StringVar(value="100")
        rows_combo = ctk.CTkOptionMenu(left_controls, variable=self.rows_per_page_var,
                                  values=["50", "100", "200", "500", "1000", "All"],
                                  command=self.on_rows_per_page_changed_optionmenu,
                                  width=100, font=FONTS["main"], dropdown_font=FONTS["main"])
        rows_combo.pack(side="left", padx=5)

        right_controls = ctk.CTkFrame(self.pagination_frame, fg_color="transparent")
        right_controls.pack(side="right", padx=10)

        ctk.CTkButton(right_controls, text="|<<", command=self.go_to_first_page, width=40, height=24, fg_color="transparent", border_width=1, text_color=("gray10", "gray90"), font=FONTS["small"]).pack(side="left", padx=2)
        ctk.CTkButton(right_controls, text="<", command=self.go_to_prev_page, width=40, height=24, fg_color="transparent", border_width=1, text_color=("gray10", "gray90"), font=FONTS["small"]).pack(side="left", padx=2)

        self.page_label = ctk.CTkLabel(right_controls, text="Page 0 of 0", width=150, font=FONTS["main"])
        self.page_label.pack(side="left", padx=5)

        ctk.CTkButton(right_controls, text=">", command=self.go_to_next_page, width=40, height=24, fg_color="transparent", border_width=1, text_color=("gray10", "gray90"), font=FONTS["small"]).pack(side="left", padx=2)
        ctk.CTkButton(right_controls, text=">>|", command=self.go_to_last_page, width=40, height=24, fg_color="transparent", border_width=1, text_color=("gray10", "gray90"), font=FONTS["small"]).pack(side="left", padx=2)

        ctk.CTkLabel(right_controls, text="Go to:", font=FONTS["main"]).pack(side="left", padx=(10, 5))
        self.page_entry = ctk.CTkEntry(right_controls, width=50, font=FONTS["main"])
        self.page_entry.pack(side="left", padx=(0, 5))
        self.page_entry.bind('<Return>', self.go_to_page)
        ctk.CTkButton(right_controls, text="Go", command=self.go_to_page, width=40, fg_color="transparent", border_width=1, text_color=("gray10", "gray90"), font=FONTS["button"]).pack(side="left")

        # Execution Mode & Parallel
        config_frame = ctk.CTkFrame(self.main_frame)
        config_frame.pack(fill="x", pady=(0, 10))

        # Col 1: Execution Sequence
        seq_col = ctk.CTkFrame(config_frame, fg_color="transparent")
        seq_col.pack(side="left", fill="both", expand=True)
        ctk.CTkLabel(seq_col, text="Execution Sequence", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)
        
        self.execution_sequence = tk.StringVar(value="full_scenario")
        ctk.CTkRadioButton(seq_col, text="Single API", variable=self.execution_sequence, value="single_api", font=FONTS["main"]).pack(anchor="w", padx=10, pady=2)
        ctk.CTkRadioButton(seq_col, text="Full Scenario", variable=self.execution_sequence, value="full_scenario", font=FONTS["main"]).pack(anchor="w", padx=10, pady=2)
        delay_row = ctk.CTkFrame(seq_col, fg_color="transparent")
        delay_row.pack(anchor="w", padx=10, pady=(4, 2))
        self.step_delay_label = ctk.CTkLabel(delay_row, text="Delay between steps:", font=FONTS["main"])
        self.step_delay_label.pack(side="left")
        self.step_delay_var = tk.StringVar(value="0")
        self.step_delay_entry = ctk.CTkEntry(delay_row, textvariable=self.step_delay_var, width=50, font=FONTS["main"])
        self.step_delay_entry.pack(side="left", padx=5)
        self.step_delay_unit_label = ctk.CTkLabel(delay_row, text="sec", font=FONTS["main"])
        self.step_delay_unit_label.pack(side="left")
        self.execution_sequence.trace_add("write", self.on_execution_sequence_change)
        self.on_execution_sequence_change()

        # Col 2: Mode
        mode_col = ctk.CTkFrame(config_frame, fg_color="transparent")
        mode_col.pack(side="left", fill="both", expand=True)
        ctk.CTkLabel(mode_col, text="Execution Mode", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)
        self.execution_mode = tk.StringVar(value=ExecutionMode.SYSTEM1_ONLY)
        ctk.CTkRadioButton(mode_col, text="System legacy (Collect Results)", variable=self.execution_mode, value=ExecutionMode.SYSTEM1_ONLY, command=self.on_execution_mode_change, font=FONTS["main"]).pack(anchor="w", padx=10, pady=2)
        ctk.CTkRadioButton(mode_col, text="System 2.0", variable=self.execution_mode, value=ExecutionMode.BOTH_SYSTEMS, command=self.on_execution_mode_change, font=FONTS["main"]).pack(anchor="w", padx=10, pady=2)

        # Col 3: Parallel
        par_col = ctk.CTkFrame(config_frame, fg_color="transparent")
        par_col.pack(side="left", fill="both", expand=True)
        ctk.CTkLabel(par_col, text="Parallel Execution", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)
        
        self.enable_parallel = tk.BooleanVar(value=False)
        ctk.CTkCheckBox(par_col, text="Enable Parallel Execution", variable=self.enable_parallel, font=FONTS["main"]).pack(anchor="w", padx=10, pady=2)
        
        w_inner = ctk.CTkFrame(par_col, fg_color="transparent")
        w_inner.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(w_inner, text="Workers:", font=FONTS["main"]).pack(side="left")
        self.worker_count = tk.StringVar(value="auto")
        ctk.CTkOptionMenu(w_inner, variable=self.worker_count, values=["auto", "2", "3", "4", "6", "8", "10", "12"], width=80, font=FONTS["main"], dropdown_font=FONTS["main"]).pack(side="left", padx=5)

        # Col 4: 2.0 Target Environment
        env_col = ctk.CTkFrame(config_frame, fg_color="transparent")
        env_col.pack(side="left", fill="both", expand=True)
        ctk.CTkLabel(env_col, text="2.0 Target Environment", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)
        
        self.new_target_env = tk.StringVar(value="test")
        self.new_target_test_radio = ctk.CTkRadioButton(env_col, text="Test", variable=self.new_target_env, value="test", font=FONTS["main"])
        self.new_target_test_radio.pack(anchor="w", padx=10, pady=2)
        self.new_target_dev_radio = ctk.CTkRadioButton(env_col, text="Dev", variable=self.new_target_env, value="dev", font=FONTS["main"])
        self.new_target_dev_radio.pack(anchor="w", padx=10, pady=2)
        self.new_target_uat_radio = ctk.CTkRadioButton(env_col, text="UAT", variable=self.new_target_env, value="uat", font=FONTS["main"])
        self.new_target_uat_radio.pack(anchor="w", padx=10, pady=2)
        self.new_target_stage_radio = ctk.CTkRadioButton(env_col, text="Stage", variable=self.new_target_env, value="stage", font=FONTS["main"])
        self.new_target_stage_radio.pack(anchor="w", padx=10, pady=2)

        # Set initial state - disabled since default is system1_only
        self.new_target_test_radio.configure(state="disabled")
        self.new_target_dev_radio.configure(state="disabled")
        self.new_target_uat_radio.configure(state="disabled")
        self.new_target_stage_radio.configure(state="disabled")

        # Specific User Login
        login_frame = ctk.CTkFrame(self.main_frame)
        login_frame.pack(fill="x", pady=(0, 10))
        
        ctk.CTkLabel(login_frame, text="Use specific user login", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)
        
        l_inner = ctk.CTkFrame(login_frame, fg_color="transparent")
        l_inner.pack(fill="x", padx=10, pady=5)
        
        self.use_specific_login = tk.BooleanVar(value=False)
        self.provider_value = tk.StringVar(value="MalaffiProvidertestD")
        self.payer_value = tk.StringVar(value="MalaffiPayertestD")
        self.tpa_value = tk.StringVar(value="MalaffiPayer2testD")
        self.pharmacy_value = tk.StringVar(value="MalaffiProvider2testD")

        def toggle_login_entry():
            state = "normal" if self.use_specific_login.get() else "disabled"
            self.provider_entry.configure(state=state)
            self.payer_entry.configure(state=state)
            self.tpa_entry.configure(state=state)
            self.pharmacy_entry.configure(state=state)

        ctk.CTkCheckBox(l_inner, text="Enable", variable=self.use_specific_login, command=toggle_login_entry, font=FONTS["main"]).pack(side="left", padx=(0, 20))

        ctk.CTkLabel(l_inner, text="Provider:", font=FONTS["main"]).pack(side="left", padx=(0, 5))
        self.provider_entry = ctk.CTkEntry(l_inner, textvariable=self.provider_value, width=150, font=FONTS["main"])
        self.provider_entry.pack(side="left", padx=5)

        ctk.CTkLabel(l_inner, text="Payer:", font=FONTS["main"]).pack(side="left", padx=(10, 5))
        self.payer_entry = ctk.CTkEntry(l_inner, textvariable=self.payer_value, width=150, font=FONTS["main"])
        self.payer_entry.pack(side="left", padx=5)

        ctk.CTkLabel(l_inner, text="TPA:", font=FONTS["main"]).pack(side="left", padx=(10, 5))
        self.tpa_entry = ctk.CTkEntry(l_inner, textvariable=self.tpa_value, width=150, font=FONTS["main"])
        self.tpa_entry.pack(side="left", padx=5)

        ctk.CTkLabel(l_inner, text="Pharmacy:", font=FONTS["main"]).pack(side="left", padx=(10, 5))
        self.pharmacy_entry = ctk.CTkEntry(l_inner, textvariable=self.pharmacy_value, width=150, font=FONTS["main"])
        self.pharmacy_entry.pack(side="left", padx=5)

        toggle_login_entry() # Initial state

        # Custom Disposition Flag
        disposition_frame = ctk.CTkFrame(self.main_frame)
        disposition_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(disposition_frame, text="Custom Disposition Flag", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)

        df_inner = ctk.CTkFrame(disposition_frame, fg_color="transparent")
        df_inner.pack(fill="x", padx=10, pady=5)

        self.use_custom_disposition_flag = tk.BooleanVar(value=False)
        self.custom_disposition_flag_value = tk.StringVar(value="")

        def toggle_disposition_entry():
            state = "normal" if self.use_custom_disposition_flag.get() else "disabled"
            self.disposition_flag_entry.configure(state=state)

        ctk.CTkCheckBox(df_inner, text="Enable", variable=self.use_custom_disposition_flag, command=toggle_disposition_entry, font=FONTS["main"]).pack(side="left", padx=(0, 20))

        ctk.CTkLabel(df_inner, text="Disposition Flag:", font=FONTS["main"]).pack(side="left", padx=(0, 5))
        self.disposition_flag_entry = ctk.CTkEntry(df_inner, textvariable=self.custom_disposition_flag_value, width=200, font=FONTS["main"])
        self.disposition_flag_entry.pack(side="left", padx=5)

        toggle_disposition_entry()  # Initial state

        # Reporting options
        reporting_frame = ctk.CTkFrame(self.main_frame)
        reporting_frame.pack(fill="x", pady=(0, 10))
        
        ctk.CTkLabel(reporting_frame, text="Allure HTML Reporting", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)
        
        r_inner = ctk.CTkFrame(reporting_frame, fg_color="transparent")
        r_inner.pack(fill="x", padx=10, pady=5)
        
        if getattr(sys, "frozen", False):
            self.allure_results_dir = tk.StringVar(value=str(Path(sys.executable).parent / "allure-results"))
        else:
            self.allure_results_dir = tk.StringVar(value="allure-results")
        ctk.CTkButton(r_inner, text="Open Last Report", command=self.open_allure_report, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left", padx=(0, 5))
        ctk.CTkButton(r_inner, text="Generate HTML", command=self.generate_allure_html_report_manual, fg_color="#1f538d", font=FONTS["button"]).pack(side="left", padx=(0, 5))
        ctk.CTkButton(r_inner, text="Save Report", command=self.save_allure_report, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left")

        # Preview table
        ctk.CTkLabel(self.main_frame, text="File Preview:", font=FONTS["sub_header"]).pack(anchor="w", padx=5)

        tree_frame = ctk.CTkFrame(self.main_frame)
        tree_frame.pack(fill="both", expand=True, pady=(0, 10))

        self.tree = ttk.Treeview(tree_frame, show='headings')
        
        v_scroll = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=self.tree.yview)
        h_scroll = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        self.tree.pack(side="left", fill="both", expand=True)
        v_scroll.pack(side="right", fill="y")
        h_scroll.pack(side="bottom", fill="x")

        # Progress and status
        progress_frame = ctk.CTkFrame(self.main_frame)
        progress_frame.pack(fill="x", pady=(0, 10))
        
        self.progress_var = tk.DoubleVar()
        self.progress_bar = ctk.CTkProgressBar(progress_frame, variable=self.progress_var)
        self.progress_bar.pack(fill="x", padx=10, pady=(10, 5))
        self.progress_bar.set(0)

        self.status_var = tk.StringVar(value="Ready")
        ctk.CTkLabel(progress_frame, textvariable=self.status_var, text_color="cyan", font=FONTS["main"]).pack(anchor="w", padx=10, pady=(0, 10))

        # Control buttons
        button_frame = ctk.CTkFrame(self.main_frame, fg_color="transparent")
        button_frame.pack(fill="x", pady=10)

        ctk.CTkButton(button_frame, text="Create Template", command=self.create_template, fg_color="#1f538d", font=FONTS["button"]).pack(side="left", padx=(0, 10))
        ctk.CTkButton(button_frame, text="Run Tests", command=self.show_execution_dialog, fg_color="green", hover_color="#006400", width=150, font=FONTS["button"]).pack(side="left", padx=(0, 10))
        ctk.CTkButton(button_frame, text="Stop", command=self.stop_tests, fg_color="red", hover_color="#8b0000", font=FONTS["button"]).pack(side="left", padx=(0, 10))
        self.regenerate_btn = ctk.CTkButton(
            button_frame,
            text="Generate Custom Report",
            command=self.regenerate_custom_report,
            fg_color="#6a0dad",
            hover_color="#4b0082",
            font=FONTS["button"]
        )
        self.regenerate_btn.pack(side="left")
        ToolTip(self.regenerate_btn, "Generate a custom report from the existing allure-results folder\n(no need to re-run tests)")

        self.ado_report_btn = ctk.CTkButton(
            button_frame, text="Report Failures ADO",
            command=self.show_ado_report_dialog,
            fg_color="#b85c00", hover_color="#7a3c00", font=FONTS["button"])
        self.ado_report_btn.pack(side="left", padx=(10, 0))
        ToolTip(self.ado_report_btn, "Report failures as bugs in Azure DevOps")

        self.ado_update_bugs_btn = ctk.CTkButton(
            button_frame, text="Update QA Bugs",
            command=self.show_update_bugs_dialog,
            fg_color="#1a6b4a", hover_color="#134d36", font=FONTS["button"])
        self.ado_update_bugs_btn.pack(side="left", padx=(10, 0))
        ToolTip(self.ado_update_bugs_btn, "For each executed test case, find its 'Ready for QA' bug in ADO\nand close it (pass) or reopen it (fail)")

        self.ado_was_ever_reported_btn = ctk.CTkButton(
            button_frame, text="Was Ever Reported",
            command=self.show_was_ever_reported_dialog,
            fg_color="#4a4a8a", hover_color="#33336b", font=FONTS["button"])
        self.ado_was_ever_reported_btn.pack(side="left", padx=(10, 0))
        ToolTip(self.ado_was_ever_reported_btn, "Find test cases that have a Closed bug in ADO\nbut no 'Ready for QA' bug — i.e. were ever reported but not currently in QA")

    # ──────────────────────────────────────────────────────────────────────────
    # Azure DevOps Reporting
    # ──────────────────────────────────────────────────────────────────────────

    def show_ado_report_dialog(self):
        """Main entry point: tag selection + preview → credentials → create bugs."""
        allure_dir = self.allure_results_dir.get() if hasattr(self, 'allure_results_dir') else "allure-results"
        if not Path(allure_dir).exists():
            messagebox.showerror("Error", f"Allure results directory not found:\n{allure_dir}")
            return

        report_tags = SYSTEM1_AVAILABLE_REPORT_TAGS
        target_tags = SYSTEM1_TARGET_TAGS

        selected = self._show_tag_selection_and_preview_dialog(allure_dir, report_tags, target_tags)
        if not selected:
            self.status_var.set("Ready")
            return

        creds = self._show_ado_credentials_dialog()
        if creds is None:
            self.status_var.set("Ready")
            return

        self._create_ado_bugs(creds, selected, target_tags)

    def _show_ado_credentials_dialog(self):
        """ADO credentials dialog with PBI. Returns dict with url, pat, project, pbi_id or None."""
        result = {}
        cancelled = [False]

        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Azure DevOps Connection")
        dialog.geometry("560x640")
        dialog.resizable(False, False)
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="Azure DevOps Wizard", font=FONTS["sub_header"]).pack(pady=(15, 5))

        url_var      = tk.StringVar(value="https://devops.malaffi.ae/ADHDS")
        pat_var      = tk.StringVar()
        project_var  = tk.StringVar()
        pbi_id_var   = tk.StringVar()
        priority_var = tk.StringVar(value="2")
        severity_var = tk.StringVar(value="3 - Medium")
        assignee_var = tk.StringVar()

        step1_frame = ctk.CTkFrame(dialog)
        step1_frame.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(step1_frame, text="1. Server Details", font=FONTS["main_bold"]).pack(anchor="w", padx=10, pady=(10, 5))

        row1 = ctk.CTkFrame(step1_frame, fg_color="transparent")
        row1.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row1, text="ADO URL:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkEntry(row1, textvariable=url_var, width=350, placeholder_text="e.g. http://server/tfs/DefaultCollection", font=FONTS["main"]).pack(side="left", padx=5)

        row3 = ctk.CTkFrame(step1_frame, fg_color="transparent")
        row3.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row3, text="PAT:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkEntry(row3, textvariable=pat_var, width=200, show="*", font=FONTS["main"]).pack(side="left", padx=5)

        status_label = ctk.CTkLabel(dialog, text="", text_color="cyan", font=FONTS["main"])
        status_label.pack(pady=5)

        step2_frame = ctk.CTkFrame(dialog)
        step2_frame.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(step2_frame, text="2. Project & PBI", font=FONTS["main_bold"]).pack(anchor="w", padx=10, pady=(10, 5))

        def on_project_selected(value):
            if value and value != "Select a project":
                pbi_entry.configure(state="normal")
                btn_verify_pbi.configure(state="normal")

        row4 = ctk.CTkFrame(step2_frame, fg_color="transparent")
        row4.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row4, text="Project:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        proj_combo = ctk.CTkOptionMenu(row4, variable=project_var, values=["Select a project"], width=250, font=FONTS["main"], state="disabled", command=on_project_selected)
        proj_combo.pack(side="left", padx=5)

        row5 = ctk.CTkFrame(step2_frame, fg_color="transparent")
        row5.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row5, text="PBI ID:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        pbi_entry = ctk.CTkEntry(row5, textvariable=pbi_id_var, width=120, font=FONTS["main"], state="disabled")
        pbi_entry.pack(side="left", padx=5)
        btn_verify_pbi = ctk.CTkButton(row5, text="Verify PBI", width=100, state="disabled")
        btn_verify_pbi.pack(side="left", padx=5)

        pbi_name_label = ctk.CTkLabel(step2_frame, text="", text_color="lightgreen", font=FONTS["main"])
        pbi_name_label.pack(anchor="w", padx=120, pady=2)

        step3_frame = ctk.CTkFrame(dialog)
        step3_frame.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(step3_frame, text="3. Bug Details", font=FONTS["main_bold"]).pack(anchor="w", padx=10, pady=(10, 5))

        row_pri = ctk.CTkFrame(step3_frame, fg_color="transparent")
        row_pri.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row_pri, text="Priority:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkOptionMenu(row_pri, variable=priority_var, values=["1", "2", "3", "4"], width=80, font=FONTS["main"]).pack(side="left", padx=5)

        row_sev = ctk.CTkFrame(step3_frame, fg_color="transparent")
        row_sev.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row_sev, text="Severity:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkOptionMenu(row_sev, variable=severity_var,
                          values=["1 - Critical", "2 - High", "3 - Medium", "4 - Low"],
                          width=200, font=FONTS["main"]).pack(side="left", padx=5)

        row_asn = ctk.CTkFrame(step3_frame, fg_color="transparent")
        row_asn.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row_asn, text="Assignee:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkEntry(row_asn, textvariable=assignee_var, width=300, font=FONTS["main"],
                     placeholder_text="Leave blank to auto-assign to yourself").pack(side="left", padx=5)

        btn_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        btn_frame.pack(fill="x", padx=20, pady=(15, 10))
        btn_cancel = ctk.CTkButton(btn_frame, text="Cancel", fg_color="gray", hover_color="#404040", font=FONTS["button"])
        btn_cancel.pack(side="left", expand=True, padx=5)
        btn_connect = ctk.CTkButton(step1_frame, text="Test Connection", width=120)
        btn_connect.pack(pady=10)
        btn_start = ctk.CTkButton(btn_frame, text="Start Reporting", fg_color="#1f538d", font=FONTS["button"], state="disabled")
        btn_start.pack(side="left", expand=True, padx=5)

        def on_test_connection():
            u, pat = url_var.get().strip(), pat_var.get().strip()
            if not u:
                status_label.configure(text="Please provide the ADO URL.", text_color="red")
                return
            status_label.configure(text="Testing connection...", text_color="cyan")
            dialog.update_idletasks()
            def do_test():
                ok, msg, projs = test_ado_connection(u, pat)
                def do_update():
                    if ok:
                        status_label.configure(text=msg, text_color="lightgreen")
                        if projs:
                            proj_combo.configure(values=projs, state="normal")
                            project_var.set(projs[0])
                            on_project_selected(projs[0])
                    else:
                        status_label.configure(text=msg, text_color="red")
                try: dialog.after(0, do_update)
                except: pass
            threading.Thread(target=do_test, daemon=True).start()

        def on_verify_pbi():
            u, pat = url_var.get().strip(), pat_var.get().strip()
            proj, pbi = project_var.get().strip(), pbi_id_var.get().strip()
            if not pbi:
                pbi_name_label.configure(text="Please enter a PBI ID.", text_color="red")
                return
            pbi_name_label.configure(text="Verifying PBI...", text_color="cyan")
            dialog.update_idletasks()
            def do_verify():
                ok, msg = get_pbi_details(u, proj, pbi, pat)
                def do_update():
                    if ok:
                        pbi_name_label.configure(text=msg, text_color="lightgreen")
                        btn_start.configure(state="normal")
                    else:
                        pbi_name_label.configure(text=msg, text_color="red")
                        btn_start.configure(state="disabled")
                try: dialog.after(0, do_update)
                except: pass
            threading.Thread(target=do_verify, daemon=True).start()

        def on_start():
            result["url"]      = url_var.get().strip()
            result["pat"]      = pat_var.get().strip()
            result["project"]  = project_var.get().strip()
            result["pbi_id"]   = pbi_id_var.get().strip()
            result["priority"] = priority_var.get().strip()
            result["severity"] = severity_var.get().strip()
            result["assignee"] = assignee_var.get().strip()
            dialog.destroy()

        def on_cancel_cmd():
            cancelled[0] = True
            dialog.destroy()

        btn_connect.configure(command=on_test_connection)
        btn_verify_pbi.configure(command=on_verify_pbi)
        btn_start.configure(command=on_start)
        btn_cancel.configure(command=on_cancel_cmd)
        dialog.protocol("WM_DELETE_WINDOW", on_cancel_cmd)
        self.frame.wait_window(dialog)

        if cancelled[0] or not result:
            return None
        return result

    def _show_tag_selection_and_preview_dialog(self, allure_dir, report_tags=None, target_tags=None):
        """Combined dialog: pick a tag → see filtered case list → preview → confirm."""
        if report_tags is None:
            report_tags = AVAILABLE_REPORT_TAGS
        if target_tags is None:
            target_tags = TARGET_TAGS
        result_ref   = [None]
        cancelled    = [False]
        failures_ref = []
        check_vars   = []
        row_frames   = []
        active_row   = [None]
        fp_groups    = {}

        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Report Failures to Azure DevOps")
        dialog.geometry("1080x660")
        dialog.minsize(900, 520)
        dialog.resizable(True, True)
        dialog.transient(self.frame)
        dialog.grab_set()

        hdr = ctk.CTkFrame(dialog, fg_color="transparent")
        hdr.pack(fill="x", padx=15, pady=(12, 0))
        ctk.CTkLabel(hdr, text="Select & Preview Failures for ADO Reporting", font=FONTS["sub_header"]).pack(side="left")

        tag_frame = ctk.CTkFrame(dialog)
        tag_frame.pack(fill="x", padx=15, pady=(8, 6))
        ctk.CTkLabel(tag_frame, text="Filter by Tag:", font=FONTS["main_bold"]).pack(side="left", padx=(10, 14), pady=8)

        selected_tag_var = tk.StringVar(value="")
        for display_name in report_tags:
            ctk.CTkRadioButton(tag_frame, text=display_name, variable=selected_tag_var, value=display_name, font=FONTS["main"]).pack(side="left", padx=10, pady=8)

        scan_status_var = tk.StringVar(value="← Select a tag to load failures")
        scan_lbl = ctk.CTkLabel(tag_frame, textvariable=scan_status_var, font=FONTS["main"], text_color="gray")
        scan_lbl.pack(side="left", padx=15)

        fp_frame = ctk.CTkFrame(dialog)
        fp_frame.pack(fill="x", padx=15, pady=(0, 4))

        content = ctk.CTkFrame(dialog, fg_color="transparent")
        content.pack(fill="both", expand=True, padx=15, pady=(0, 5))

        left = ctk.CTkFrame(content)
        left.pack(side="left", fill="both", expand=True, padx=(0, 6))

        ctrl_row = ctk.CTkFrame(left, fg_color="transparent")
        ctrl_row.pack(fill="x", padx=5, pady=(5, 2))
        count_lbl = ctk.CTkLabel(ctrl_row, text="No failures loaded", font=FONTS["main"], text_color="cyan")
        count_lbl.pack(side="left", padx=5)

        def _select_all():
            for v in check_vars: v.set(True)
            unique_only_active[0] = False
            dup_options_frame.pack_forget()
            dup_comments_var.set(False)
        def _select_none():
            for v in check_vars: v.set(False)
            unique_only_active[0] = False
            dup_options_frame.pack_forget()
            dup_comments_var.set(False)
        def _select_unique_only():
            keep = set()
            for indices in fp_groups.values():
                keep.add(indices[0])
            for i, v in enumerate(check_vars):
                v.set(i in keep)
            unique_only_active[0] = True
            dup_options_frame.pack(fill="x", padx=5, pady=(0, 2))

        ctk.CTkButton(ctrl_row, text="None", width=55, fg_color="gray", hover_color="#404040", font=FONTS["small"], command=_select_none).pack(side="right", padx=2)
        ctk.CTkButton(ctrl_row, text="All", width=55, fg_color="gray", hover_color="#404040", font=FONTS["small"], command=_select_all).pack(side="right", padx=2)
        ctk.CTkButton(ctrl_row, text="Unique Only", width=90, fg_color="#1f538d", hover_color="#153a6b", font=FONTS["small"], command=_select_unique_only).pack(side="right", padx=2)

        # ── Duplicate-comments option (shown only when "Unique Only" is active)
        dup_comments_var = tk.BooleanVar(value=False)
        dup_max_var      = tk.StringVar(value="3")
        dup_prefix_var   = tk.StringVar(value="")
        dup_options_frame = ctk.CTkFrame(left, fg_color="transparent")
        # not packed until _select_unique_only() is called
        dup_row1 = ctk.CTkFrame(dup_options_frame, fg_color="transparent")
        dup_row1.pack(fill="x")
        ctk.CTkCheckBox(
            dup_row1, text="Add duplicates in comments",
            variable=dup_comments_var, font=FONTS["small"],
        ).pack(side="left", padx=(5, 10))
        ctk.CTkLabel(dup_row1, text="Max:", font=FONTS["small"]).pack(side="left")
        tk.Spinbox(
            dup_row1, from_=1, to=10, textvariable=dup_max_var,
            width=3, justify="center",
            bg="#2b2b2b", fg="white", buttonbackground="#3a3a3a",
            relief="flat", highlightthickness=0,
        ).pack(side="left", padx=(2, 4))
        ctk.CTkLabel(
            dup_row1, text="comments per bug", font=FONTS["small"],
        ).pack(side="left")
        dup_row2 = ctk.CTkFrame(dup_options_frame, fg_color="transparent")
        dup_row2.pack(fill="x", pady=(3, 0))
        ctk.CTkLabel(dup_row2, text="Comment prefix text:", font=FONTS["small"]).pack(side="left", padx=(5, 4))
        ctk.CTkEntry(dup_row2, textvariable=dup_prefix_var, font=FONTS["small"], width=260).pack(side="left")

        list_scroll = ctk.CTkScrollableFrame(left)
        list_scroll.pack(fill="both", expand=True, padx=5, pady=(0, 5))

        right = ctk.CTkFrame(content, width=440)
        right.pack(side="right", fill="both")
        right.pack_propagate(False)

        ctk.CTkLabel(right, text="Bug Preview", font=FONTS["main_bold"]).pack(anchor="w", padx=10, pady=(8, 2))
        preview_title_var = tk.StringVar(value="")
        ctk.CTkLabel(right, textvariable=preview_title_var, font=FONTS["main"], wraplength=410, justify="left", anchor="w").pack(fill="x", padx=10, pady=(0, 4))

        att_sel_frame = ctk.CTkFrame(right, fg_color="transparent")
        att_sel_frame.pack(fill="x", padx=10, pady=(2, 2))
        row1 = ctk.CTkFrame(att_sel_frame, fg_color="transparent")
        row1.pack(fill="x", pady=1)

        def _on_att_check_changed():
            if active_row[0]:
                try:
                    idx = row_frames.index(active_row[0])
                    failure = failures_ref[idx]
                    if "include_attachments" not in failure:
                        failure["include_attachments"] = {}
                    failure["include_attachments"]["send"]               = inc_send_var.get()
                    failure["include_attachments"]["include_assertions"] = inc_assertions_var.get()
                except ValueError:
                    pass

        def _apply_attachments_to_all():
            for f in failures_ref:
                if "include_attachments" not in f:
                    f["include_attachments"] = {}
                f["include_attachments"]["send"]               = inc_send_var.get()
                f["include_attachments"]["include_assertions"] = inc_assertions_var.get()

        inc_send_var       = tk.BooleanVar(value=True)
        inc_assertions_var = tk.BooleanVar(value=True)

        ctk.CTkLabel(row1, text="Include:", font=FONTS["main_bold"]).pack(side="left", padx=(0, 5))
        ctk.CTkCheckBox(row1, text="Send",       variable=inc_send_var,       font=FONTS["small"], width=10, command=_on_att_check_changed).pack(side="left", padx=5)
        ctk.CTkCheckBox(row1, text="Assertions", variable=inc_assertions_var, font=FONTS["small"], width=10, command=_on_att_check_changed).pack(side="left", padx=5)
        ctk.CTkButton(row1, text="Apply All", width=60, height=20, font=FONTS["small"], command=_apply_attachments_to_all).pack(side="right", padx=(5, 0))

        ctk.CTkFrame(right, height=1, fg_color="gray30").pack(fill="x", padx=8, pady=2)

        tab_view = ctk.CTkTabview(right)
        tab_view.pack(fill="both", expand=True, padx=5, pady=(0, 5))
        for t in ["Send", "Assertions"]:
            tab_view.add(t)

        no_sel_lbl = ctk.CTkLabel(right, text="← Click a case\nto preview the bug", font=FONTS["main"], text_color="gray50", justify="center")
        no_sel_lbl.place(relx=0.5, rely=0.55, anchor="center")

        def _make_tab_txt(tab_name):
            frame = tab_view.tab(tab_name)
            txt = tk.Text(frame, bg="#1a1a2e", fg="#d4d4d4", font=("Courier", 9), wrap="word", relief="flat", state="disabled", borderwidth=0)
            sc = ttk.Scrollbar(frame, command=txt.yview)
            txt.configure(yscrollcommand=sc.set)
            sc.pack(side="right", fill="y")
            txt.pack(side="left", fill="both", expand=True)
            return txt

        txt_send       = _make_tab_txt("Send")
        txt_assertions = _make_tab_txt("Assertions")

        def _fill(widget, content):
            widget.configure(state="normal")
            widget.delete("1.0", tk.END)
            widget.insert("1.0", content if content else "(no content)")
            widget.configure(state="disabled")

        def _show_preview(failure):
            no_sel_lbl.place_forget()
            atts = failure["attachments"]
            preview_title_var.set(f"Title: {failure['test_case_name']}\nRule ID: {failure['rule_id']}")
            _fill(txt_send, atts.get("send", ""))
            _ASSERT_LABELS = [
                ("assert_2ffe8e", "Check Against Transaction (Not Found)"),
                ("assert_8b450b", "Check Against Transaction (Found)"),
                ("assert_d43b5a", "Transaction FileID Found"),
                ("assert_568d49", "Check Error Message"),
                ("assert_c67011", "Compare Template & Downloaded File"),
                ("assert_c67045", "Check Set Downloaded"),
                ("assert_c67048", "Check Error Code"),
                ("assert_c63496", "Reconciliation Check"),
                ("assert_e4c52a", "Element Value Verification"),
                ("assert_f41a3b", "File Found Check"),
                ("assert_f17d8e", "File Not Found Check"),
            ]
            parts = []
            for key, label in _ASSERT_LABELS:
                c = atts.get(key, "")
                if c:
                    parts.append(f"{'='*60}\n{label}\n{'='*60}\n{c}")
            _fill(txt_assertions, "\n\n".join(parts))
            inc = failure.get("include_attachments", {})
            inc_send_var.set(inc.get("send", True))
            inc_assertions_var.set(inc.get("include_assertions", True))

        bot = ctk.CTkFrame(dialog, fg_color="transparent")
        bot.pack(fill="x", padx=15, pady=(2, 12))

        def on_cancel():
            cancelled[0] = True
            dialog.destroy()

        tags_var = tk.StringVar(value="")
        check_existing_var = tk.BooleanVar(value=True)

        # Fingerprint criteria BooleanVars
        fp_use_transaction_type = tk.BooleanVar(value=True)
        fp_use_rule_id          = tk.BooleanVar(value=True)
        fp_use_scenario_type    = tk.BooleanVar(value=True)
        fp_use_object           = tk.BooleanVar(value=True)
        fp_use_element          = tk.BooleanVar(value=True)
        fp_use_failure_reason   = tk.BooleanVar(value=True)
        fp_use_error_text       = tk.BooleanVar(value=False)
        fp_use_bug_title        = tk.BooleanVar(value=False)
        unique_only_active      = [False]

        def _fingerprint(failure):
            tags = [t.lower() for t in failure.get("all_tags", [])]
            parts = []
            for prefix, var in (
                ("transaction type - ", fp_use_transaction_type),
                ("ruleid - ",           fp_use_rule_id),
                ("scenario type - ",    fp_use_scenario_type),
                ("object - ",           fp_use_object),
                ("element - ",          fp_use_element),
            ):
                if var.get():
                    val = next((t for t in tags if t.startswith(prefix)), "")
                    parts.append(val)
            if fp_use_failure_reason.get():
                reasons = tuple(sorted(t for t in tags if t in target_tags))
                parts.append(str(reasons))
            if fp_use_error_text.get():
                parts.append(failure.get("error_text", ""))
            if fp_use_bug_title.get():
                parts.append(failure.get("test_case_name", ""))
            return tuple(parts)

        def _recompute_fingerprints(*_):
            if not failures_ref:
                return
            fp_groups.clear()
            for i, f in enumerate(failures_ref):
                fp = _fingerprint(f)
                fp_groups.setdefault(fp, []).append(i)
            distinct_bugs = len(fp_groups)
            scan_status_var.set(
                f"{len(failures_ref)} failure(s) found  |  {distinct_bugs} distinct bug(s)"
                f"  — use \"Unique Only\" to auto-select one per bug"
            )
            scan_lbl.configure(text_color="lightgreen")
            _rebuild_list(list(failures_ref))
            _select_unique_only()

        def on_report():
            custom_tags = [t.strip() for t in tags_var.get().split(",") if t.strip()]
            check_existing = check_existing_var.get()
            fp_criteria = {
                "transaction_type": fp_use_transaction_type.get(),
                "rule_id":          fp_use_rule_id.get(),
                "scenario_type":    fp_use_scenario_type.get(),
                "object":           fp_use_object.get(),
                "element":          fp_use_element.get(),
                "failure_reason":   fp_use_failure_reason.get(),
                "error_text":       fp_use_error_text.get(),
                "bug_title":        fp_use_bug_title.get(),
            }
            selected_pairs = [(i, failures_ref[i]) for i, v in enumerate(check_vars) if v.get()]
            selected = [f for _, f in selected_pairs]
            if dup_comments_var.get():
                try:
                    max_dup = min(10, max(1, int(dup_max_var.get())))
                except ValueError:
                    max_dup = 3
                for orig_idx, f in selected_pairs:
                    fp = _fingerprint(f)
                    group = fp_groups.get(fp, [orig_idx])
                    f["duplicate_comments"] = [
                        failures_ref[j] for j in group if j != orig_idx
                    ][:max_dup]
                    f["duplicate_comment_prefix"] = dup_prefix_var.get().strip()
            for f in selected:
                f["custom_tags"] = custom_tags
                f["check_existing"] = check_existing
                f["fp_criteria"] = fp_criteria
            result_ref[0] = selected
            dialog.destroy()

        ctk.CTkButton(bot, text="Cancel", command=on_cancel, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left")
        ctk.CTkLabel(bot, text="Add Tags (comma separated):", font=FONTS["main_bold"]).pack(side="left", padx=(20, 5))
        ctk.CTkEntry(bot, textvariable=tags_var, font=FONTS["main"], width=250, placeholder_text="e.g. Sprint-1, Regression").pack(side="left", padx=5)
        ctk.CTkCheckBox(bot, text="Check Existing Bugs", variable=check_existing_var, font=FONTS["small"]).pack(side="left", padx=15)

        ctk.CTkLabel(fp_frame, text="Unique By:", font=FONTS["main_bold"]).pack(side="left", padx=(10, 8), pady=6)
        for _fp_var, _fp_lbl in [
            (fp_use_transaction_type, "Transaction Type"),
            (fp_use_rule_id,          "Rule ID"),
            (fp_use_scenario_type,    "Scenario Type"),
            (fp_use_object,           "Object"),
            (fp_use_element,          "Element"),
            (fp_use_failure_reason,   "Failure Reason"),
            (fp_use_error_text,       "Error Text"),
            (fp_use_bug_title,        "Bug Title"),
        ]:
            ctk.CTkCheckBox(fp_frame, text=_fp_lbl, variable=_fp_var, font=FONTS["main"],
                            command=_recompute_fingerprints).pack(side="left", padx=8, pady=6)

        report_btn = ctk.CTkButton(bot, text="Report Selected (0)", command=on_report, fg_color="#b85c00", hover_color="#7a3c00", font=FONTS["button"], state="disabled")
        report_btn.pack(side="right")

        def _update_count(*_):
            n = sum(1 for v in check_vars if v.get())
            count_lbl.configure(text=f"{n} case(s) selected")
            report_btn.configure(text=f"Report Selected ({n})", state="normal" if n > 0 else "disabled")

        def _rebuild_list(failures):
            for w in list_scroll.winfo_children():
                w.destroy()
            check_vars.clear()
            row_frames.clear()
            failures_ref.clear()
            failures_ref.extend(failures)
            active_row[0] = None
            no_sel_lbl.place(relx=0.5, rely=0.55, anchor="center")
            preview_title_var.set("")
            for tw in (txt_send, txt_assertions):
                _fill(tw, "")
            if not failures:
                ctk.CTkLabel(list_scroll, text="No failures found for this tag.", font=FONTS["main"], text_color="gray").pack(pady=20)
                count_lbl.configure(text="0 cases selected")
                report_btn.configure(text="Report Selected (0)", state="disabled")
                return
            for i, failure in enumerate(failures):
                var = tk.BooleanVar(value=True)
                var.trace_add("write", _update_count)
                check_vars.append(var)
                base_color = "#2a2a2a" if i % 2 == 0 else "#242424"
                row = ctk.CTkFrame(list_scroll, fg_color=base_color, corner_radius=4)
                row.pack(fill="x", pady=1, padx=2)
                row_frames.append(row)
                ctk.CTkCheckBox(row, text="", variable=var, width=24).pack(side="left", padx=(6, 2), pady=4)
                lbl = ctk.CTkLabel(row, text=f"[{failure['rule_id']}]  {failure['test_case_name']}", anchor="w", font=FONTS["main"], wraplength=270)
                lbl.pack(side="left", padx=(2, 8), pady=4, fill="x", expand=True)
                def _on_click(ev=None, f=failure, r=row, idx=i):
                    if active_row[0] and active_row[0] is not r:
                        prev = active_row[0]
                        try:
                            pi = row_frames.index(prev)
                            prev.configure(fg_color="#2a2a2a" if pi % 2 == 0 else "#242424")
                        except ValueError:
                            pass
                    active_row[0] = r
                    r.configure(fg_color="#1f538d")
                    _show_preview(f)
                row.bind("<Button-1>", _on_click)
                lbl.bind("<Button-1>", _on_click)
            _update_count()

        def _on_tag_changed(*_):
            tag = selected_tag_var.get()
            if not tag:
                return
            scan_status_var.set("Scanning…")
            scan_lbl.configure(text_color="cyan")
            dialog.update_idletasks()

            def _do_scan():
                try:
                    found = scan_assertion_failures_by_tag(allure_dir, tag, report_tags)
                    try:
                        dialog.after(0, lambda: _on_done(found))
                    except tk.TclError:
                        pass
                except Exception as exc:
                    try:
                        dialog.after(0, lambda: _on_err(str(exc)))
                    except tk.TclError:
                        pass

            def _on_done(found):
                for f in found:
                    if "include_attachments" not in f:
                        f["include_attachments"] = {"send": True, "include_assertions": True}
                n = len(found)
                fp_groups.clear()
                for i, f in enumerate(found):
                    fp = _fingerprint(f)
                    fp_groups.setdefault(fp, []).append(i)
                distinct_bugs = len(fp_groups)
                if n:
                    scan_status_var.set(f"{n} failure(s) found  |  {distinct_bugs} distinct bug(s)  — use \"Unique Only\" to auto-select one per bug")
                else:
                    scan_status_var.set("No failures found")
                scan_lbl.configure(text_color="lightgreen" if n else "orange")
                _rebuild_list(found)

            def _on_err(err):
                scan_status_var.set(f"Error: {err}")
                scan_lbl.configure(text_color="red")

            threading.Thread(target=_do_scan, daemon=True).start()

        selected_tag_var.trace_add("write", _on_tag_changed)
        dialog.protocol("WM_DELETE_WINDOW", on_cancel)
        self.frame.wait_window(dialog)

        if cancelled[0]:
            return None
        return result_ref[0]

    def _create_ado_bugs(self, creds, failures, target_tags=None):
        """Create ADO bugs for every selected failure and show a progress dialog."""
        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Reporting Bugs to Azure DevOps")
        dialog.geometry("620x450")
        dialog.resizable(True, True)
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="Creating Bugs in Azure DevOps...", font=FONTS["sub_header"]).pack(pady=(15, 5))
        progress_var = tk.DoubleVar()
        progress_bar = ctk.CTkProgressBar(dialog, variable=progress_var)
        progress_bar.pack(fill="x", padx=20, pady=5)
        progress_bar.set(0)
        status_lbl = ctk.CTkLabel(dialog, text="Starting...", font=FONTS["main"], text_color="cyan")
        status_lbl.pack(pady=(0, 5))
        log_box = tk.Text(dialog, height=15, bg="#1e1e1e", fg="white", font=("Courier", 10), state="disabled", relief="flat")
        log_box.pack(fill="both", expand=True, padx=20, pady=5)
        log_box.tag_config("ok",    foreground="#4caf50")
        log_box.tag_config("error", foreground="#f44336")
        log_box.tag_config("info",  foreground="#82aaff")
        close_btn = ctk.CTkButton(dialog, text="Close", command=dialog.destroy, fg_color="gray", hover_color="#404040", font=FONTS["button"], state="disabled")
        close_btn.pack(pady=(5, 15))

        def log(text, tag="info"):
            log_box.configure(state="normal")
            log_box.insert(tk.END, text + "\n", tag)
            log_box.see(tk.END)
            log_box.configure(state="disabled")
            dialog.update_idletasks()

        def run_reporting():
            total   = len(failures)
            success = 0
            errors  = 0
            skipped = 0
            for i, failure in enumerate(failures):
                failure["priority"] = creds.get("priority") or None
                failure["severity"] = creds.get("severity") or None
                failure["assignee"] = creds.get("assignee") or None
                tc_name = failure["test_case_name"]
                status_lbl.configure(text=f"Reporting {i+1}/{total}: {tc_name[:60]}...")
                log(f"[{i+1}/{total}] Processing: {tc_name}", "info")
                try:
                    skip = False
                    if failure.get("check_existing"):
                        log("  Searching for existing bugs in ADO...", "info")
                        exists, skip_msg = check_existing_bug(
                            url=creds["url"], project=creds["project"],
                            pat=creds["pat"], failure=failure, target_tags=target_tags,
                            fp_criteria=failure.get("fp_criteria"))
                        if exists:
                            log(f"  ⏭ {skip_msg}", "ok")
                            skipped += 1
                            skip = True
                    if not skip:
                        log("  Creating bug...", "info")
                        ok, msg, wi_id = create_ado_bug(
                            url=creds["url"], project=creds["project"],
                            pat=creds["pat"], pbi_id=creds["pbi_id"], failure=failure)
                        if ok:
                            success += 1
                            log(f"  ✔ {msg}", "ok")
                            dup_failures = failure.get("duplicate_comments", [])
                            dup_prefix = failure.get("duplicate_comment_prefix", "")
                            for k, dup in enumerate(dup_failures):
                                try:
                                    c_ok, c_msg = add_failure_comment(
                                        url=creds["url"],
                                        project=creds["project"],
                                        pat=creds["pat"],
                                        work_item_id=wi_id,
                                        failure=dup,
                                        prefix_text=dup_prefix,
                                    )
                                    if c_ok:
                                        log(f"    💬 Duplicate {k+1}/{len(dup_failures)}: comment added", "ok")
                                    else:
                                        log(f"    ⚠ Duplicate {k+1} comment failed: {c_msg}", "error")
                                except Exception as ce:
                                    log(f"    ⚠ Duplicate {k+1} comment error: {ce}", "error")
                        else:
                            errors += 1
                            log(f"  ✘ {msg}", "error")
                except Exception as e:
                    errors += 1
                    log(f"  ✘ Unexpected error: {e}", "error")
                progress_var.set((i + 1) / total)
                dialog.update_idletasks()
            status_lbl.configure(text=f"Done — {success} created, {skipped} skipped, {errors} error(s).")
            log(f"\n═══ Finished: {success} created, {skipped} skipped, {errors} errors. ═══", "ok" if errors == 0 else "error")
            close_btn.configure(state="normal")
            self.status_var.set(f"ADO Reporting complete: {success} created, {skipped} skipped.")

        threading.Thread(target=run_reporting, daemon=True).start()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)
        self.frame.wait_window(dialog)

    # ──────────────────────────────────────────────────────────────────────────
    # Azure DevOps Update QA Bugs
    # ──────────────────────────────────────────────────────────────────────────

    def show_update_bugs_dialog(self):
        """Entry point for the 'Update QA Bugs' flow: match config → credentials → run."""
        allure_dir = self.allure_results_dir.get() if hasattr(self, 'allure_results_dir') else "allure-results"
        if not Path(allure_dir).exists():
            messagebox.showerror("Error", f"Allure results directory not found:\n{allure_dir}")
            return
        test_cases = scan_all_assertion_results(allure_dir)
        if not test_cases:
            messagebox.showinfo("No Results", "No test case results found in the allure-results directory.")
            return
        match_fields = self._show_bug_match_config_dialog()
        if match_fields is None:
            self.status_var.set("Ready")
            return
        creds = self._show_ado_connection_dialog()
        if creds is None:
            self.status_var.set("Ready")
            return
        self._run_update_bugs(creds, test_cases, match_fields)

    def _show_bug_match_config_dialog(self):
        """Checkboxes for which fields to use when matching bugs. Returns a set of field keys or None."""
        result_ref = [None]

        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Bug Match Configuration")
        dialog.geometry("460x600")
        dialog.resizable(False, False)
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="Match bugs using:", font=FONTS["sub_header"]).pack(pady=(15, 5))
        ctk.CTkLabel(dialog, text="Select the fields to use when searching for a matching\n'Ready for QA' bug in ADO for each test case.", font=FONTS["main"], justify="center").pack(pady=(0, 10))

        fields = [
            ("transaction_type", "Transaction Type"),
            ("rule_id",          "Rule ID"),
            ("scenario_type",    "Scenario Type"),
            ("object",           "Object"),
            ("element",          "Element"),
            ("failure_tag",      "Failure Tag"),
            ("bug_title",        "Bug Title"),
        ]
        tooltips = {
            "failure_tag": (
                "Matches the ADO bug by the failure type tag on the test case.\n"
                "e.g. \"System 2 Failure\", \"System 1 Failure\",\n"
                "     \"System 2 Error Text Failure\", \"Dual System Error\"\n\n"
                "Note: this condition is automatically skipped for passed test\n"
                "cases since they carry no failure tag."
            ),
            "bug_title": (
                "Matches the ADO bug title exactly against the test case name.\n"
                "Bug titles are stored in the format:\n"
                "  \"Test Case: {Transaction Type} - {Description}\"\n"
                "e.g. \"Test Case: Get New Transactions - Verify Person Register ACK appears\""
            ),
        }

        vars_ = {}
        checks_frame = ctk.CTkFrame(dialog)
        checks_frame.pack(fill="x", padx=30, pady=5)
        for key, label in fields:
            var = tk.BooleanVar(value=True)
            vars_[key] = var
            cb = ctk.CTkCheckBox(checks_frame, text=label, variable=var, font=FONTS["main"])
            cb.pack(anchor="w", pady=4)
            if key in tooltips:
                ToolTip(cb, tooltips[key])

        # Dynamic description label
        ctk.CTkLabel(dialog, text="Matching logic:", font=FONTS["main"], text_color="#aaaaaa").pack(pady=(10, 2))
        desc_label = ctk.CTkLabel(
            dialog,
            text="",
            font=FONTS["main"],
            wraplength=400,
            justify="left",
            text_color="#e0e0e0",
        )
        desc_label.pack(padx=30, anchor="w")

        def build_description():
            selected = {k for k, v in vars_.items() if v.get()}
            if not selected:
                return "No fields selected — no bugs will match."
            criteria = []
            for key, _ in fields:
                if key not in selected:
                    continue
                if key == "transaction_type":
                    criteria.append("Transaction Type tag matches")
                elif key == "rule_id":
                    criteria.append("Rule ID tag matches")
                elif key == "scenario_type":
                    criteria.append("Scenario Type tag matches")
                elif key == "object":
                    criteria.append("Object tag matches")
                elif key == "element":
                    criteria.append("Element tag matches")
                elif key == "failure_tag":
                    criteria.append("Failure Tag matches (skipped for passed tests)")
                elif key == "bug_title":
                    criteria.append("bug title = test case name (case-insensitive)")
            if len(criteria) == 1:
                joined = criteria[0]
            elif len(criteria) == 2:
                joined = f"{criteria[0]} AND {criteria[1]}"
            else:
                joined = ", ".join(criteria[:-1]) + f", and {criteria[-1]}"
            return f"A bug matches a test case when: {joined}."

        def update_description(*_):
            desc_label.configure(text=build_description())

        for var in vars_.values():
            var.trace_add("write", update_description)

        update_description()

        def on_confirm():
            selected = {k for k, v in vars_.items() if v.get()}
            if not selected:
                messagebox.showwarning("No Fields Selected", "Select at least one field to match on.", parent=dialog)
                return
            result_ref[0] = selected
            dialog.destroy()

        def on_cancel():
            dialog.destroy()

        btn_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        btn_frame.pack(pady=(10, 15))
        ctk.CTkButton(btn_frame, text="Continue", command=on_confirm, font=FONTS["button"]).pack(side="left", padx=10)
        ctk.CTkButton(btn_frame, text="Cancel", command=on_cancel, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left", padx=10)

        self.frame.wait_window(dialog)
        return result_ref[0]

    def _show_ado_connection_dialog(self):
        """Lightweight ADO connection dialog (URL, PAT, project — no PBI). Returns dict or None."""
        result = {}
        cancelled = [False]

        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Azure DevOps Connection")
        dialog.geometry("560x420")
        dialog.resizable(False, False)
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="Azure DevOps Connection", font=FONTS["sub_header"]).pack(pady=(15, 5))

        url_var     = tk.StringVar(value="https://devops.malaffi.ae/ADHDS")
        pat_var     = tk.StringVar()
        project_var = tk.StringVar()

        step1_frame = ctk.CTkFrame(dialog)
        step1_frame.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(step1_frame, text="1. Server Details", font=FONTS["main_bold"]).pack(anchor="w", padx=10, pady=(10, 5))

        row1 = ctk.CTkFrame(step1_frame, fg_color="transparent")
        row1.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row1, text="ADO URL:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkEntry(row1, textvariable=url_var, width=350, placeholder_text="e.g. http://server/tfs/DefaultCollection", font=FONTS["main"]).pack(side="left", padx=5)

        row2 = ctk.CTkFrame(step1_frame, fg_color="transparent")
        row2.pack(fill="x", padx=10, pady=2)
        ctk.CTkLabel(row2, text="PAT:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        ctk.CTkEntry(row2, textvariable=pat_var, width=200, show="*", font=FONTS["main"]).pack(side="left", padx=5)

        btn_connect = ctk.CTkButton(step1_frame, text="Test Connection", width=120)
        btn_connect.pack(pady=10)

        status_label = ctk.CTkLabel(dialog, text="", text_color="cyan", font=FONTS["main"])
        status_label.pack(pady=5)

        step2_frame = ctk.CTkFrame(dialog)
        step2_frame.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(step2_frame, text="2. Project", font=FONTS["main_bold"]).pack(anchor="w", padx=10, pady=(10, 5))

        row3 = ctk.CTkFrame(step2_frame, fg_color="transparent")
        row3.pack(fill="x", padx=10, pady=(5, 15))
        ctk.CTkLabel(row3, text="Project:", width=100, anchor="w", font=FONTS["main"]).pack(side="left")
        proj_combo = ctk.CTkOptionMenu(row3, variable=project_var, values=["Select a project"], width=250, font=FONTS["main"], state="disabled")
        proj_combo.pack(side="left", padx=5)

        btn_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        btn_frame.pack(fill="x", padx=20, pady=(15, 10))
        btn_cancel = ctk.CTkButton(btn_frame, text="Cancel", fg_color="gray", hover_color="#404040", font=FONTS["button"])
        btn_cancel.pack(side="left", expand=True, padx=5)
        btn_start = ctk.CTkButton(btn_frame, text="Continue", fg_color="#1f538d", font=FONTS["button"], state="disabled")
        btn_start.pack(side="left", expand=True, padx=5)

        def on_test_connection():
            u, pat = url_var.get().strip(), pat_var.get().strip()
            if not u:
                status_label.configure(text="Please provide the ADO URL.", text_color="red")
                return
            status_label.configure(text="Testing connection...", text_color="cyan")
            dialog.update_idletasks()
            def do_test():
                ok, msg, projs = test_ado_connection(u, pat)
                def do_update():
                    if ok:
                        status_label.configure(text=msg, text_color="lightgreen")
                        if projs:
                            proj_combo.configure(values=projs, state="normal")
                            project_var.set(projs[0])
                            btn_start.configure(state="normal")
                    else:
                        status_label.configure(text=msg, text_color="red")
                try: dialog.after(0, do_update)
                except: pass
            threading.Thread(target=do_test, daemon=True).start()

        def on_start():
            result["url"]     = url_var.get().strip()
            result["pat"]     = pat_var.get().strip()
            result["project"] = project_var.get().strip()
            dialog.destroy()

        def on_cancel_cmd():
            cancelled[0] = True
            dialog.destroy()

        btn_connect.configure(command=on_test_connection)
        btn_start.configure(command=on_start)
        btn_cancel.configure(command=on_cancel_cmd)
        dialog.protocol("WM_DELETE_WINDOW", on_cancel_cmd)
        self.frame.wait_window(dialog)

        if cancelled[0] or not result:
            return None
        return result

    def _run_update_bugs(self, creds, test_cases, match_fields=None):
        """For each test case, find its 'Ready for QA' bug and close or reopen it."""
        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Updating QA Bugs in Azure DevOps")
        dialog.geometry("650x480")
        dialog.resizable(True, True)
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="Updating 'Ready for QA' Bugs...", font=FONTS["sub_header"]).pack(pady=(15, 5))
        progress_var = tk.DoubleVar()
        progress_bar = ctk.CTkProgressBar(dialog, variable=progress_var)
        progress_bar.pack(fill="x", padx=20, pady=5)
        progress_bar.set(0)
        status_lbl = ctk.CTkLabel(dialog, text="Starting...", font=FONTS["main"], text_color="cyan")
        status_lbl.pack(pady=(0, 5))
        log_box = tk.Text(dialog, height=15, bg="#1e1e1e", fg="white", font=("Courier", 10), state="disabled", relief="flat")
        log_box.pack(fill="both", expand=True, padx=20, pady=5)
        log_box.tag_config("ok",      foreground="#4caf50")
        log_box.tag_config("error",   foreground="#f44336")
        log_box.tag_config("info",    foreground="#82aaff")
        log_box.tag_config("skipped", foreground="#aaaaaa")
        close_btn = ctk.CTkButton(dialog, text="Close", command=dialog.destroy, fg_color="gray", hover_color="#404040", font=FONTS["button"], state="disabled")
        close_btn.pack(pady=(5, 15))

        def log(text, tag="info"):
            log_box.configure(state="normal")
            log_box.insert(tk.END, text + "\n", tag)
            log_box.see(tk.END)
            log_box.configure(state="disabled")
            dialog.update_idletasks()

        def run():
            log("Fetching all 'Ready for QA' bugs from ADO...", "info")
            status_lbl.configure(text="Fetching 'Ready for QA' bugs from ADO...")
            dialog.update_idletasks()

            bugs, fetch_err = get_all_ready_for_qa_bugs(
                url=creds["url"], project=creds["project"], pat=creds["pat"]
            )

            if fetch_err:
                log(f"✘ Failed to fetch bugs: {fetch_err}", "error")
                close_btn.configure(state="normal")
                return

            if not bugs:
                log("No 'Ready for QA' bugs found in ADO.", "skipped")
                close_btn.configure(state="normal")
                return

            log(f"Found {len(bugs)} 'Ready for QA' bug(s). Matching against {len(test_cases)} executed test cases...", "info")

            total    = len(bugs)
            closed   = 0
            reopened = 0
            skipped  = 0
            errors   = 0

            for i, bug in enumerate(bugs):
                wi_id = bug["id"]
                title = bug["title"]
                status_lbl.configure(text=f"Processing {i+1}/{total}: Bug #{wi_id}...")
                log(f"[{i+1}/{total}] Bug #{wi_id}: {title[:60]}", "info")
                try:
                    tc = match_bug_to_test_case(bug, test_cases, match_fields)

                    if tc is None:
                        log("  ⏭ No matching executed test case found", "skipped")
                        skipped += 1
                        progress_var.set((i + 1) / total)
                        dialog.update_idletasks()
                        continue

                    tc_name = tc["test_case_name"]
                    status  = tc["status"]
                    log(f"  Matched: {tc_name}  [{status}]", "info")

                    new_state = "Ready for Release" if tc.get("qa_outcome") == "close" else "Reopened"
                    ok, msg = update_bug_for_qa_result(
                        url=creds["url"], project=creds["project"],
                        pat=creds["pat"], work_item_id=wi_id,
                        test_case=tc, new_state=new_state)
                    if ok:
                        if new_state == "Ready for Release":
                            closed += 1
                        else:
                            reopened += 1
                        log(f"  ✔ {msg}", "ok")
                    else:
                        errors += 1
                        log(f"  ✘ {msg}", "error")
                except Exception as e:
                    errors += 1
                    log(f"  ✘ Unexpected error: {e}", "error")
                progress_var.set((i + 1) / total)
                dialog.update_idletasks()

            summary = f"Done — {closed} closed, {reopened} reopened, {skipped} skipped (no matching test case), {errors} error(s)."
            status_lbl.configure(text=summary)
            log(f"\n═══ {summary} ═══", "ok" if errors == 0 else "error")
            close_btn.configure(state="normal")
            self.status_var.set(f"QA bug update complete: {closed} closed, {reopened} reopened.")

        threading.Thread(target=run, daemon=True).start()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)
        self.frame.wait_window(dialog)

    def show_was_ever_reported_dialog(self):
        """Entry point for 'Was Ever Reported': match config → credentials → run."""
        allure_dir = self.allure_results_dir.get() if hasattr(self, 'allure_results_dir') else "allure-results"
        if not Path(allure_dir).exists():
            messagebox.showerror("Error", f"Allure results directory not found:\n{allure_dir}")
            return
        test_cases = scan_all_results(allure_dir)
        if not test_cases:
            messagebox.showinfo("No Results", "No test case results found in the allure-results directory.")
            return
        match_fields = self._show_bug_match_config_dialog()
        if match_fields is None:
            self.status_var.set("Ready")
            return
        creds = self._show_ado_connection_dialog()
        if creds is None:
            self.status_var.set("Ready")
            return
        self._run_was_ever_reported(creds, test_cases, match_fields)

    def _run_was_ever_reported(self, creds, test_cases, match_fields=None):
        """For each Closed bug that matches a test case in the current run and has no
        matching 'Ready for QA' bug, report it as 'was ever reported'."""
        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Was Ever Reported — Azure DevOps")
        dialog.geometry("650x520")
        dialog.resizable(True, True)
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="Was Ever Reported", font=FONTS["sub_header"]).pack(pady=(15, 5))
        progress_var = tk.DoubleVar()
        progress_bar = ctk.CTkProgressBar(dialog, variable=progress_var)
        progress_bar.pack(fill="x", padx=20, pady=5)
        progress_bar.set(0)
        status_lbl = ctk.CTkLabel(dialog, text="Starting...", font=FONTS["main"], text_color="cyan")
        status_lbl.pack(pady=(0, 5))
        log_box = tk.Text(dialog, height=15, bg="#1e1e1e", fg="white", font=("Courier", 10), state="disabled", relief="flat")
        log_box.pack(fill="both", expand=True, padx=20, pady=5)
        log_box.tag_config("ok",      foreground="#4caf50")
        log_box.tag_config("error",   foreground="#f44336")
        log_box.tag_config("info",    foreground="#82aaff")
        log_box.tag_config("skipped", foreground="#aaaaaa")

        btn_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        btn_frame.pack(pady=(5, 15))

        # matches keyed by test_case_name: {tc, bugs: []}
        matches_by_tc = {}

        def export_to_excel():
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            save_path = filedialog.asksaveasfilename(
                title="Save Was Ever Reported",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile="was_ever_reported.xlsx",
            )
            if not save_path:
                return
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Was Ever Reported"

                headers = ["Test Case Name", "All Tags", "Closed Bug IDs", "Closed Bug Titles"]
                header_fill = PatternFill("solid", fgColor="1F4E79")
                header_font = Font(bold=True, color="FFFFFF")
                for col, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(wrap_text=True, vertical="center")

                for row_idx, (tc_name, entry) in enumerate(matches_by_tc.items(), 2):
                    tc      = entry["tc"]
                    bugs    = entry["bugs"]
                    all_tags_str   = "; ".join(tc.get("all_tags", []))
                    bug_ids_str    = ", ".join(f"#{b['id']}" for b in bugs)
                    bug_titles_str = " | ".join(b["title"] for b in bugs)

                    ws.cell(row=row_idx, column=1, value=tc_name)
                    ws.cell(row=row_idx, column=2, value=all_tags_str)
                    ws.cell(row=row_idx, column=3, value=bug_ids_str)
                    ws.cell(row=row_idx, column=4, value=bug_titles_str)
                    for col in range(1, 5):
                        ws.cell(row=row_idx, column=col).alignment = Alignment(wrap_text=True, vertical="top")

                ws.column_dimensions["A"].width = 55
                ws.column_dimensions["B"].width = 60
                ws.column_dimensions["C"].width = 20
                ws.column_dimensions["D"].width = 60

                wb.save(save_path)
                messagebox.showinfo("Exported", f"Report saved to:\n{save_path}", parent=dialog)
            except Exception as e:
                messagebox.showerror("Export Failed", str(e), parent=dialog)

        export_btn = ctk.CTkButton(btn_frame, text="Export to Excel", command=export_to_excel,
                                   fg_color="#1a6b4a", hover_color="#134d36",
                                   font=FONTS["button"], state="disabled")
        export_btn.pack(side="left", padx=10)

        close_btn = ctk.CTkButton(btn_frame, text="Close", command=dialog.destroy,
                                  fg_color="gray", hover_color="#404040",
                                  font=FONTS["button"], state="disabled")
        close_btn.pack(side="left", padx=10)

        def log(text, tag="info"):
            log_box.configure(state="normal")
            log_box.insert(tk.END, text + "\n", tag)
            log_box.see(tk.END)
            log_box.configure(state="disabled")
            dialog.update_idletasks()

        def run():
            log("Fetching all Closed bugs from ADO...", "info")
            status_lbl.configure(text="Fetching Closed bugs from ADO...")
            dialog.update_idletasks()

            closed_bugs, err1 = get_all_closed_bugs(
                url=creds["url"], project=creds["project"], pat=creds["pat"]
            )
            if err1:
                log(f"✘ Failed to fetch Closed bugs: {err1}", "error")
                close_btn.configure(state="normal")
                return

            if not closed_bugs:
                log("No Closed bugs found in ADO.", "skipped")
                close_btn.configure(state="normal")
                return

            log(f"Found {len(closed_bugs)} Closed bug(s). Fetching 'Ready for QA' bugs...", "info")
            status_lbl.configure(text="Fetching 'Ready for QA' bugs from ADO...")
            dialog.update_idletasks()

            rfqa_bugs, err2 = get_all_ready_for_qa_bugs(
                url=creds["url"], project=creds["project"], pat=creds["pat"]
            )
            if err2:
                log(f"✘ Failed to fetch 'Ready for QA' bugs: {err2}", "error")
                close_btn.configure(state="normal")
                return

            log(f"Found {len(rfqa_bugs)} 'Ready for QA' bug(s). Matching against {len(test_cases)} executed test cases...", "info")

            total    = len(closed_bugs)
            reported = 0
            skipped  = 0
            errors   = 0

            for i, bug in enumerate(closed_bugs):
                wi_id = bug["id"]
                title = bug["title"]
                status_lbl.configure(text=f"Processing {i+1}/{total}: Bug #{wi_id}...")
                try:
                    tc = match_bug_to_test_case(bug, test_cases, match_fields)

                    if tc is None:
                        skipped += 1
                        progress_var.set((i + 1) / total)
                        dialog.update_idletasks()
                        continue

                    has_rfqa = any(
                        match_bug_to_test_case(rfqa_bug, [tc], match_fields) is not None
                        for rfqa_bug in rfqa_bugs
                    )

                    if has_rfqa:
                        skipped += 1
                        progress_var.set((i + 1) / total)
                        dialog.update_idletasks()
                        continue

                    reported += 1
                    tc_name = tc["test_case_name"]
                    status  = tc["status"]
                    log(f"[{reported}] Bug #{wi_id}: {title[:60]}", "ok")
                    log(f"     Matched: {tc_name}  [{status}]", "info")

                    if tc_name not in matches_by_tc:
                        matches_by_tc[tc_name] = {"tc": tc, "bugs": []}
                    matches_by_tc[tc_name]["bugs"].append(bug)

                except Exception as e:
                    errors += 1
                    log(f"  ✘ Unexpected error on Bug #{wi_id}: {e}", "error")

                progress_var.set((i + 1) / total)
                dialog.update_idletasks()

            summary = (
                f"Done — {reported} was-ever-reported, "
                f"{skipped} skipped (no test case match or has RFQA bug), "
                f"{errors} error(s)."
            )
            status_lbl.configure(text=summary)
            log(f"\n═══ {summary} ═══", "ok" if errors == 0 else "error")
            if matches_by_tc:
                export_btn.configure(state="normal")
            close_btn.configure(state="normal")
            self.status_var.set(f"Was Ever Reported complete: {reported} found.")

        threading.Thread(target=run, daemon=True).start()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)
        self.frame.wait_window(dialog)

    def on_sheet_selected_optionmenu(self, selected):
        if selected:
            self.current_sheet = selected
            self.status_var.set(f"Sheet '{selected}' selected. Click Load to preview.")

    def on_execution_sequence_change(self, *args):
        # Step delay only applies between transactions of a Full Scenario chain
        enabled = self.execution_sequence.get() == "full_scenario"
        self.step_delay_entry.configure(state="normal" if enabled else "disabled")
        text_color = ("gray10", "gray90") if enabled else ("gray60", "gray45")
        self.step_delay_label.configure(text_color=text_color)
        self.step_delay_unit_label.configure(text_color=text_color)

    def on_execution_mode_change(self):
        """Enable/disable 2.0 Target Environment options based on execution mode"""
        if self.execution_mode.get() == ExecutionMode.BOTH_SYSTEMS:  # System 2.0 selected
            self.new_target_test_radio.configure(state="normal")
            self.new_target_dev_radio.configure(state="normal")
            self.new_target_uat_radio.configure(state="normal")
            self.new_target_stage_radio.configure(state="normal")
        else:  # System legacy selected
            self.new_target_test_radio.configure(state="disabled")
            self.new_target_dev_radio.configure(state="disabled")
            self.new_target_uat_radio.configure(state="disabled")
            self.new_target_stage_radio.configure(state="disabled")

    def on_rows_per_page_changed_optionmenu(self, choice):
        self.rows_per_page_var.set(choice)
        self.on_rows_per_page_changed()

    def browse_file(self):
        file_path = filedialog.askopenfilename(title="Select Excel File", filetypes=[("Excel files", "*.xlsx *.xls")])
        if file_path:
            self.file_path_var.set(file_path)
            self.all_sheets_data = {}
            self.detect_sheets()

    def detect_sheets(self):
        file_path = self.file_path_var.get()
        if not file_path: return
        try:
            self.sheet_names = self.excel_handler.get_sheet_names(file_path)
            if len(self.sheet_names) > 1:
                self.sheet_combo.configure(values=self.sheet_names)
                self.sheet_combo.set(self.sheet_names[0])
                self.current_sheet = self.sheet_names[0]
                self.sheet_frame.pack(fill="x", pady=(0, 10), after=self.main_frame.winfo_children()[1])
                self.status_var.set(f"File has {len(self.sheet_names)} sheets. Select a sheet and click Load.")
            else:
                self.sheet_frame.pack_forget()
                self.sheet_combo.configure(values=[])
                self.current_sheet = self.sheet_names[0] if self.sheet_names else None
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read file sheets: {str(e)}")

    def refresh_sheet(self):
        if self.current_sheet:
            self.load_sheet_data(self.current_sheet)

    def load_file(self):
        file_path = self.file_path_var.get()
        if not file_path:
            messagebox.showwarning("Warning", "Please select a file first")
            return
        if not self.sheet_names:
            self.detect_sheets()
        sheet_to_load = self.current_sheet if self.current_sheet else (self.sheet_names[0] if self.sheet_names else None)
        if sheet_to_load:
            self.load_sheet_data(sheet_to_load)
        else:
            messagebox.showerror("Error", "No sheets found in the Excel file")

    def show_execution_dialog(self):
        file_path = self.file_path_var.get()
        if not file_path or not self.sheet_names:
            messagebox.showwarning("Warning", "Please load a file first")
            return
        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Test Execution Options")
        dialog.geometry("550x780")
        dialog.transient(self.frame)
        dialog.grab_set()

        main_container = ctk.CTkScrollableFrame(dialog)
        main_container.pack(fill="both", expand=True, padx=20, pady=20)

        ctk.CTkLabel(main_container, text="Select Sheets to Test", font=FONTS["title"]).pack(pady=(0, 10))
        ctk.CTkLabel(main_container, text=f"File: {os.path.basename(file_path)}", text_color="gray").pack(pady=(0, 5))
        ctk.CTkLabel(main_container, text=f"Total sheets: {len(self.sheet_names)}", text_color="gray").pack(pady=(0, 20))

        scope_var = tk.StringVar(value="current")

        # Placeholders for widgets referenced in on_scope_change
        case_listbox = None
        load_cases_btn = None

        def on_scope_change():
            if case_listbox and load_cases_btn:
                if scope_var.get() == "specific":
                    case_listbox.configure(state="normal")
                    load_cases_btn.configure(state="normal")
                else:
                    case_listbox.configure(state="disabled")
                    load_cases_btn.configure(state="disabled")

        current_frame = ctk.CTkFrame(main_container, border_width=1)
        current_frame.pack(fill="x", pady=(0, 10))
        ctk.CTkRadioButton(current_frame, text="Current Sheet Only", variable=scope_var, value="current", font=FONTS["main_bold"], command=on_scope_change).pack(anchor="w", padx=10, pady=5)
        ctk.CTkLabel(current_frame, text=f"   → Test only: {self.current_sheet}", text_color="#1f538d", font=FONTS["main"]).pack(anchor="w", padx=20, pady=(0, 5))

        all_frame = ctk.CTkFrame(main_container, border_width=1)
        all_frame.pack(fill="x", pady=(0, 10))
        ctk.CTkRadioButton(all_frame, text="All Sheets in File", variable=scope_var, value="all", font=FONTS["main_bold"], command=on_scope_change).pack(anchor="w", padx=10, pady=5)
        ctk.CTkLabel(all_frame, text=f"   → Test all {len(self.sheet_names)} sheets", text_color="#1f538d", font=FONTS["main"]).pack(anchor="w", padx=20, pady=(0, 5))

        selected_frame = ctk.CTkFrame(main_container, border_width=1)
        selected_frame.pack(fill="x", pady=(0, 10), expand=True)
        ctk.CTkRadioButton(selected_frame, text="Selected Sheets", variable=scope_var, value="selected", font=FONTS["main_bold"], command=on_scope_change).pack(anchor="w", padx=10, pady=5)
        ctk.CTkLabel(selected_frame, text="   → Choose specific sheets below:", text_color="#1f538d", font=FONTS["main"]).pack(anchor="w", padx=20, pady=(0, 5))

        listbox_frame = ctk.CTkFrame(selected_frame, height=150)
        listbox_frame.pack(fill="both", expand=True, padx=20, pady=5)
        
        scrollbar = ttk.Scrollbar(listbox_frame)
        scrollbar.pack(side="right", fill="y")
        sheet_listbox = tk.Listbox(listbox_frame, selectmode=tk.MULTIPLE, yscrollcommand=scrollbar.set, height=8, bg="#333333", fg="white", selectbackground="#1f538d", font=("Arial", 12))
        sheet_listbox.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=sheet_listbox.yview)

        for sheet in self.sheet_names:
            sheet_listbox.insert(tk.END, f"{sheet}")

        quick_btn_frame = ctk.CTkFrame(selected_frame, fg_color="transparent")
        quick_btn_frame.pack(fill="x", padx=20, pady=5)
        ctk.CTkButton(quick_btn_frame, text="Select All", command=lambda: sheet_listbox.selection_set(0, tk.END), width=80, fg_color="transparent", border_width=1, font=FONTS["button"]).pack(side="left", padx=(0, 5))
        ctk.CTkButton(quick_btn_frame, text="Clear All", command=lambda: sheet_listbox.selection_clear(0, tk.END), width=80, fg_color="transparent", border_width=1, font=FONTS["button"]).pack(side="left")

        # ── Specific Cases Option ──────────────────────────────────────────────
        specific_frame = ctk.CTkFrame(main_container, border_width=1)
        specific_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkRadioButton(specific_frame, text="Specific Cases", variable=scope_var, value="specific", font=FONTS["main_bold"], command=on_scope_change).pack(anchor="w", padx=10, pady=5)
        ctk.CTkLabel(specific_frame, text="   → Choose specific test cases below:", text_color="#1f538d", font=FONTS["main"]).pack(anchor="w", padx=20, pady=(0, 5))

        # Frame for sheet selection for cases
        case_sheet_frame = ctk.CTkFrame(specific_frame, fg_color="transparent")
        case_sheet_frame.pack(fill="x", padx=20, pady=5)

        ctk.CTkLabel(case_sheet_frame, text="Load from:", font=FONTS["main"]).pack(side="left", padx=(0, 5))
        case_sheet_var = tk.StringVar(value=self.current_sheet if self.current_sheet else self.sheet_names[0])
        case_sheet_combo = ctk.CTkComboBox(case_sheet_frame, values=self.sheet_names, variable=case_sheet_var, width=150, font=FONTS["main"])
        case_sheet_combo.pack(side="left", padx=(0, 10))

        test_cases_data = {}

        def load_test_cases():
            selected_sheet = case_sheet_var.get()
            try:
                if selected_sheet not in self.all_sheets_data:
                    data = self.excel_handler.load_excel(file_path, sheet_name=selected_sheet)
                    self.all_sheets_data[selected_sheet] = data
                else:
                    data = self.all_sheets_data[selected_sheet]

                case_listbox.configure(state="normal")
                case_listbox.delete(0, tk.END)
                test_cases_data.clear()

                for row in data:
                    tc_id = row.get('TC ID', '')
                    if tc_id:
                        case_listbox.insert(tk.END, str(tc_id))
                        test_cases_data[str(tc_id)] = row

                case_count_label.configure(text=f"   {len(test_cases_data)} cases loaded from '{selected_sheet}'")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to load test cases: {str(e)}", parent=dialog)

        load_cases_btn = ctk.CTkButton(case_sheet_frame, text="Load Cases", command=load_test_cases, width=100, font=FONTS["button"])
        load_cases_btn.pack(side="left")

        case_count_label = ctk.CTkLabel(specific_frame, text="   Click 'Load Cases' to see available test cases", text_color="gray", font=FONTS["main"])
        case_count_label.pack(anchor="w", padx=20, pady=(0, 5))

        case_listbox_frame = ctk.CTkFrame(specific_frame, height=150)
        case_listbox_frame.pack(fill="both", expand=True, padx=20, pady=5)

        case_scrollbar = ttk.Scrollbar(case_listbox_frame)
        case_scrollbar.pack(side="right", fill="y")
        case_listbox = tk.Listbox(case_listbox_frame, selectmode=tk.EXTENDED, yscrollcommand=case_scrollbar.set, height=8, bg="#333333", fg="white", selectbackground="#1f538d", font=("Arial", 11), state="disabled")
        case_listbox.pack(side="left", fill="both", expand=True)
        case_scrollbar.config(command=case_listbox.yview)

        case_quick_btn_frame = ctk.CTkFrame(specific_frame, fg_color="transparent")
        case_quick_btn_frame.pack(fill="x", padx=20, pady=5)

        ctk.CTkButton(case_quick_btn_frame, text="Select All", command=lambda: case_listbox.selection_set(0, tk.END), width=80, fg_color="transparent", border_width=1, font=FONTS["button"]).pack(side="left", padx=(0, 5))
        ctk.CTkButton(case_quick_btn_frame, text="Clear All", command=lambda: case_listbox.selection_clear(0, tk.END), width=80, fg_color="transparent", border_width=1, font=FONTS["button"]).pack(side="left")


        button_frame = ctk.CTkFrame(main_container, fg_color="transparent")
        button_frame.pack(fill="x", pady=20)

        def run():
            scope = scope_var.get()
            specific_cases = None
            if scope == "current":
                sheets_to_test = [self.current_sheet]
            elif scope == "all":
                sheets_to_test = self.sheet_names
            elif scope == "selected":
                selected_indices = sheet_listbox.curselection()
                if not selected_indices:
                    messagebox.showwarning("No Selection", "Please select at least one sheet", parent=dialog)
                    return
                sheets_to_test = [self.sheet_names[i] for i in selected_indices]
            elif scope == "specific":
                selected_indices = case_listbox.curselection()
                if not selected_indices:
                    messagebox.showwarning("No Selection", "Please select at least one test case", parent=dialog)
                    return
                specific_cases = [case_listbox.get(i) for i in selected_indices]
                sheets_to_test = [case_sheet_var.get()]
            else:
                sheets_to_test = [self.current_sheet]

            dialog.destroy()
            self.execute_tests(sheets_to_test, specific_cases=specific_cases)

        ctk.CTkButton(button_frame, text="Cancel", command=dialog.destroy, fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="right", padx=(5, 0))
        ctk.CTkButton(button_frame, text="Run Tests", command=run, fg_color="green", hover_color="#006400", font=FONTS["button"]).pack(side="right")

    def execute_tests(self, sheets_to_test, specific_cases=None):
        file_path = self.file_path_var.get()
        for sheet in sheets_to_test:
            if sheet not in self.all_sheets_data:
                try:
                    data = self.excel_handler.load_excel(file_path, sheet_name=sheet)
                    self.all_sheets_data[sheet] = data
                except Exception as e:
                    messagebox.showerror("Error", f"Failed to load sheet '{sheet}': {str(e)}")
                    return

        # Count total (filtered if needed)
        if specific_cases:
            total_tests = sum(
                1 for sheet in sheets_to_test
                for row in self.all_sheets_data[sheet]
                if str(row.get('TC ID', '')) in specific_cases
            )
            confirm_msg = f"Ready to execute {total_tests} specific test case(s)?"
        else:
            total_tests = sum(len(self.all_sheets_data[sheet]) for sheet in sheets_to_test)
            confirm_msg = f"Ready to execute {total_tests} tests across {len(sheets_to_test)} sheets?"

        if not messagebox.askyesno("Confirm Execution", confirm_msg):
            return

        if messagebox.askyesno("Custom Report", "Do you want to generate a custom report?"):
            self.show_custom_report_dialog()
            if not self.generate_custom_report:
                return
        else:
            self.generate_custom_report = False
            self.custom_report_sheets = []

        self.status_var.set(f"Executing {total_tests} tests...")

        result, message = validate_assertion_excel(self.all_sheets_data)
        if result:
            if not message:
                self.run_allure_tests_only(file_path, ", ".join(sheets_to_test), specific_cases)
            else:
                messagebox.showwarning("Warning", f"Extra columns: {message}")
                self.run_allure_tests_only(file_path, ", ".join(sheets_to_test), specific_cases)
        else:
            messagebox.showerror("Error", f"Missing columns: {message}")

    def run_allure_tests_only(self, file_path, sheets_to_test, specific_cases=None):
        try:
            os.environ[EnvVar.SOAP_EXCEL_FILE] = file_path
            os.environ[EnvVar.EXCEL_FILE_TEST_SHEETS] = sheets_to_test
            os.environ[EnvVar.SOAP_EXECUTION_MODE] = self.execution_mode.get()
            os.environ[EnvVar.SOAP_EXECUTION_SEQUENCE] = self.execution_sequence.get()
            step_delay = 0.0
            if self.execution_sequence.get() == "full_scenario":
                try:
                    step_delay = max(0.0, float(self.step_delay_var.get() or "0"))
                except ValueError:
                    step_delay = 0.0
            os.environ[EnvVar.SOAP_STEP_DELAY] = str(step_delay)
            os.environ[EnvVar.TARGET_ENVIRONMENT] = "pte"
            os.environ[EnvVar.NEW_TARGET_ENVIRONMENT] = self.new_target_env.get()

            # Set specific cases filter if provided
            if specific_cases:
                os.environ[EnvVar.SPECIFIC_TEST_CASES] = ",".join(specific_cases)
            else:
                if EnvVar.SPECIFIC_TEST_CASES in os.environ:
                    del os.environ[EnvVar.SPECIFIC_TEST_CASES]

            # Clear Rule ID filter (not used in assertion tab)
            if EnvVar.SPECIFIC_RULE_IDS in os.environ:
                del os.environ[EnvVar.SPECIFIC_RULE_IDS]

            if self.use_specific_login.get():
                os.environ[EnvVar.USE_SPECIFIC_LOGIN] = "true"
                os.environ[EnvVar.PROVIDER_VALUE] = self.provider_value.get()
                os.environ[EnvVar.PAYER_VALUE] = self.payer_value.get()
                os.environ[EnvVar.TPA_VALUE] = self.tpa_value.get()
                os.environ[EnvVar.PHARMACY_VALUE] = self.pharmacy_value.get()
            else:
                for key in (EnvVar.USE_SPECIFIC_LOGIN, EnvVar.PROVIDER_VALUE, EnvVar.PAYER_VALUE, EnvVar.TPA_VALUE, EnvVar.PHARMACY_VALUE):
                    if key in os.environ:
                        del os.environ[key]

            if self.use_custom_disposition_flag.get():
                os.environ[EnvVar.CUSTOM_DISPOSITION_FLAG] = self.custom_disposition_flag_value.get()
            else:
                if EnvVar.CUSTOM_DISPOSITION_FLAG in os.environ:
                    del os.environ[EnvVar.CUSTOM_DISPOSITION_FLAG]

            allure_dir = self.allure_results_dir.get()
            self._clear_allure_results(allure_dir)
            os.environ[EnvVar.DEPENDENCY_STATE_FILE] = str(Path(allure_dir).resolve() / '_dep_state.json')

            if getattr(sys, 'frozen', False):
                project_root = Path(sys.executable).parent
                python_exe = project_root / "python_runtime" / "python.exe"
            else:
                current_file = Path(__file__).resolve()
                project_root = current_file.parent.parent
                python_exe = sys.executable

            wrapper_path = project_root / "utils" / "assertion_allure_wrapper.py"

            cmd = [str(python_exe), "-m", "pytest", str(wrapper_path), "--alluredir", allure_dir, "--tb=short", "-v"]

            if self.enable_parallel.get():
                worker_count = self.worker_count.get()
                if worker_count == "auto":
                    cmd.extend(["-n", "auto"])
                else:
                    cmd.extend(["-n", worker_count])

            def run_pytest():
                try:
                    self.current_process = subprocess.Popen(
                        cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, stdin=subprocess.DEVNULL,
                        text=True, cwd=str(project_root), creationflags=subprocess.CREATE_NO_WINDOW
                    )
                    stdout, stderr = self.current_process.communicate()
                    result_code = self.current_process.returncode

                    if result_code == 0 or result_code == 1:
                        self.status_var.set("Allure tests completed successfully")
                        if self.generate_custom_report and self.custom_report_sheets:
                            try:
                                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                                default_name = f"custom_report_{timestamp}.xlsx"
                                report_path = filedialog.asksaveasfilename(
                                    title="Save Custom Report",
                                    defaultextension=".xlsx",
                                    initialfile=default_name,
                                    filetypes=[("Excel files", "*.xlsx")]
                                )
                                if report_path:
                                    os.environ[EnvVar.SOAP_EXECUTION_MODE] = ExecutionMode.SYSTEM1_ONLY
                                    generate_custom_combined_report(
                                        allure_dir,
                                        report_path,
                                        self.custom_report_sheets,
                                        assert_error_text_enabled=False
                                    )
                                    self.last_custom_report_path = report_path
                                    self.last_allure_dir = allure_dir
                                    messagebox.showinfo("Excel Generated", f"Custom report generated at:\n{report_path}")
                                else:
                                    self.status_var.set("Custom report generation cancelled")
                            except Exception as e:
                                messagebox.showerror("Excel Error", f"Failed to generate Custom Report: {e}")
                    else:
                        self.status_var.set(f"Tests completed with issues (exit code: {result_code})")

                    self.show_pytest_output(stdout, stderr)

                except Exception as e:
                    self.status_var.set(f"Failed to run Allure tests: {str(e)}")
                    messagebox.showerror("Execution Error", str(e))
                finally:
                    self.current_process = None
                    self.testing_active = False

            thread = threading.Thread(target=run_pytest)
            thread.daemon = True
            thread.start()

        except Exception as e:
            messagebox.showerror("Error", f"Failed to start Allure tests: {str(e)}")

    def show_custom_report_dialog(self):
        """Configure sheets and columns for the custom Excel report."""
        # Assertion testing always tags failures as "System 1 Failure" only —
        # there is no System 2 or comparison logic — so always use system1_only.
        available_columns = get_available_rules_summary_columns(ExecutionMode.SYSTEM1_ONLY, assert_error_text_enabled=False)

        s1_name = "Legacy System"
        is_system1_mode = True

        default_criteria = [f'{s1_name} Failure']

        editing_config = [{'name': 'Rules Summary', 'columns': list(available_columns), 'criteria': default_criteria}]

        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Custom Report Configuration")
        dialog.geometry("800x700")
        dialog.minsize(700, 600)
        dialog.transient(self.frame)
        dialog.grab_set()
        dialog.resizable(True, True)

        btn_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        btn_frame.pack(fill="x", side="bottom", padx=10, pady=10)

        main_scroll_container = ctk.CTkScrollableFrame(dialog)
        main_scroll_container.pack(fill="both", expand=True, padx=5, pady=(5, 0))

        # === Sheet management ===
        top_frame = ctk.CTkFrame(main_scroll_container)
        top_frame.pack(fill="x", padx=5, pady=(5, 5))
        ctk.CTkLabel(top_frame, text="Report Sheets", font=FONTS["sub_header"]).pack(anchor="w", padx=10, pady=5)

        sheet_list_frame = ctk.CTkFrame(top_frame)
        sheet_list_frame.pack(fill="x", padx=10, pady=5)

        sheet_listbox = tk.Listbox(sheet_list_frame, height=5, bg="#333333", fg="white",
                                   selectbackground="#1f538d", highlightthickness=0, bd=0,
                                   font=("Arial", 11), exportselection=False)
        sheet_listbox.pack(side="left", fill="both", expand=True)

        sheet_controls = ctk.CTkFrame(sheet_list_frame, fg_color="transparent")
        sheet_controls.pack(side="right", fill="y", padx=5)

        # === Criteria + Columns ===
        bottom_frame = ctk.CTkFrame(main_scroll_container)
        bottom_frame.pack(fill="both", expand=True, padx=5, pady=(5, 5))

        criteria_frame = ctk.CTkFrame(bottom_frame)
        criteria_frame.pack(fill="both", expand=True, padx=5, pady=5)
        ctk.CTkLabel(criteria_frame, text="Monitoring Criteria (Relationship: OR)", font=FONTS["main_bold"]).pack(anchor="w", padx=5, pady=(5, 0))

        all_criteria_names = [
            f'{s1_name} Failure',
            'Failed - No Failure Tag',
            'No Failure',
            'Default Statistics',
            'Transactions Default Statistics',
        ]

        _EXCLUSIVE_STATS = {'Default Statistics', 'Transactions Default Statistics'}
        self.criteria_vars = {name: tk.BooleanVar() for name in all_criteria_names}
        self._criteria_checkboxes = {}

        criteria_scroll_frame = ctk.CTkScrollableFrame(criteria_frame, height=120, fg_color="transparent")
        criteria_scroll_frame.pack(fill="both", expand=True, padx=5, pady=5)

        def get_current_sheet_idx_ref():
            sel = sheet_listbox.curselection()
            return sel[0] if sel else None

        def set_columns_locked(locked):
            state = "disabled" if locked else "normal"
            listbox_avail.configure(state=state)
            listbox_sel.configure(state=state)
            for btn in col_buttons:
                btn.configure(state=state)

        def on_criteria_change(name):
            idx = get_current_sheet_idx_ref()
            if idx is None:
                return
            sheet_conf = editing_config[idx]
            current = sheet_conf.get('criteria', [])

            if name in _EXCLUSIVE_STATS:
                if self.criteria_vars[name].get():
                    for other, var in self.criteria_vars.items():
                        if other != name:
                            var.set(False)
                            if other in current:
                                current.remove(other)
                        cb = self._criteria_checkboxes.get(other)
                        if cb:
                            cb.configure(state="disabled")
                    if name not in current:
                        current.append(name)
                    set_columns_locked(True)
                else:
                    for other, var in self.criteria_vars.items():
                        if other == name:
                            continue
                        cb = self._criteria_checkboxes.get(other)
                        if cb:
                            cb.configure(state="normal")
                    if name in current:
                        current.remove(name)
                    set_columns_locked(False)
            else:
                for stat_name in _EXCLUSIVE_STATS:
                    if self.criteria_vars[stat_name].get():
                        self.criteria_vars[stat_name].set(False)
                        if stat_name in current:
                            current.remove(stat_name)
                for other in self.criteria_vars:
                    if other in _EXCLUSIVE_STATS:
                        continue
                    cb = self._criteria_checkboxes.get(other)
                    if cb:
                        cb.configure(state="normal")
                set_columns_locked(False)

                if self.criteria_vars[name].get():
                    if name not in current:
                        current.append(name)
                else:
                    if name in current:
                        current.remove(name)

            sheet_conf['criteria'] = current

        for name, var in self.criteria_vars.items():
            if name in ('Failed - No Failure Tag', 'Default Statistics'):
                sep = ctk.CTkLabel(criteria_scroll_frame, text="─" * 30, text_color="gray", font=("Arial", 9))
                sep.pack(anchor="w", padx=10)
            cb = ctk.CTkCheckBox(criteria_scroll_frame, text=name, variable=var,
                                 command=lambda n=name: on_criteria_change(n), font=FONTS["main"])
            self._criteria_checkboxes[name] = cb
            cb.pack(anchor="w", padx=10, pady=5)

        # === Columns ===
        current_sheet_label = ctk.CTkLabel(bottom_frame, text="Columns Configuration", font=FONTS["sub_header"])
        current_sheet_label.pack(anchor="w", padx=10, pady=5)

        cols_container = ctk.CTkFrame(bottom_frame, fg_color="transparent")
        cols_container.pack(fill="both", expand=True, padx=5, pady=5)

        left_frame = ctk.CTkFrame(cols_container)
        left_frame.pack(side="left", fill="both", expand=True, padx=5)
        ctk.CTkLabel(left_frame, text="Available Columns", font=FONTS["main_bold"]).pack()
        listbox_avail = tk.Listbox(left_frame, selectmode=tk.MULTIPLE, bg="#333333", fg="white",
                                   highlightthickness=0, bd=0, font=("Arial", 10), exportselection=False)
        listbox_avail.pack(side="left", fill="both", expand=True, padx=5, pady=5)

        mid_frame = ctk.CTkFrame(cols_container, fg_color="transparent")
        mid_frame.pack(side="left", fill="y", padx=5)

        right_frame = ctk.CTkFrame(cols_container)
        right_frame.pack(side="left", fill="both", expand=True, padx=5)
        ctk.CTkLabel(right_frame, text="Selected Columns", font=FONTS["main_bold"]).pack()
        listbox_sel = tk.Listbox(right_frame, selectmode=tk.MULTIPLE, bg="#333333", fg="white",
                                 highlightthickness=0, bd=0, font=("Arial", 10), exportselection=False)
        listbox_sel.pack(side="left", fill="both", expand=True, padx=5, pady=5)

        def get_current_sheet_idx():
            sel = sheet_listbox.curselection()
            return sel[0] if sel else None

        def refresh_criteria_view():
            idx = get_current_sheet_idx()
            if idx is None:
                for var in self.criteria_vars.values():
                    var.set(False)
                return
            sheet_conf = editing_config[idx]
            current = sheet_conf.get('criteria', [])
            for name, var in self.criteria_vars.items():
                var.set(name in current)
            is_exclusive = any(s in current for s in _EXCLUSIVE_STATS)
            for name, cb in self._criteria_checkboxes.items():
                if name in _EXCLUSIVE_STATS and name in current:
                    continue
                cb.configure(state="disabled" if is_exclusive else "normal")
            set_columns_locked(is_exclusive)

        def refresh_columns_view():
            idx = get_current_sheet_idx()
            if idx is None:
                current_sheet_label.configure(text="No sheet selected")
                listbox_avail.delete(0, tk.END)
                listbox_sel.delete(0, tk.END)
                return
            sheet_conf = editing_config[idx]
            current_sheet_label.configure(text=f"Columns for '{sheet_conf['name']}'")
            selected_cols = sheet_conf['columns']
            listbox_avail.delete(0, tk.END)
            for col in available_columns:
                if col not in selected_cols:
                    listbox_avail.insert(tk.END, col)
            listbox_sel.delete(0, tk.END)
            for col in selected_cols:
                listbox_sel.insert(tk.END, col)
            refresh_criteria_view()

        sheet_listbox.bind('<<ListboxSelect>>', lambda e: refresh_columns_view())

        def refresh_sheet_list():
            sheet_listbox.delete(0, tk.END)
            for sheet in editing_config:
                sheet_listbox.insert(tk.END, sheet['name'])
            if editing_config:
                sheet_listbox.select_set(0)
                refresh_columns_view()

        def add_sheet():
            new_name = simpledialog.askstring("New Sheet", "Enter sheet name:", parent=dialog)
            if new_name:
                count = 1
                base_name = new_name
                existing = [s['name'] for s in editing_config]
                while new_name in existing:
                    new_name = f"{base_name}_{count}"
                    count += 1
                editing_config.append({'name': new_name, 'columns': list(available_columns), 'criteria': []})
                refresh_sheet_list()
                sheet_listbox.select_clear(0, tk.END)
                sheet_listbox.select_set(len(editing_config) - 1)
                refresh_columns_view()

        def remove_sheet():
            sel = sheet_listbox.curselection()
            if not sel:
                return
            if len(editing_config) <= 1:
                messagebox.showwarning("Warning", "Report must have at least one sheet.", parent=dialog)
                return
            del editing_config[sel[0]]
            refresh_sheet_list()

        def rename_sheet():
            sel = sheet_listbox.curselection()
            if not sel:
                return
            idx = sel[0]
            new_name = simpledialog.askstring("Rename Sheet", "Enter new sheet name:",
                                              initialvalue=editing_config[idx]['name'], parent=dialog)
            if new_name and new_name != editing_config[idx]['name']:
                existing = [s['name'] for i, s in enumerate(editing_config) if i != idx]
                if new_name in existing:
                    messagebox.showerror("Error", "Sheet name already exists.", parent=dialog)
                    return
                editing_config[idx]['name'] = new_name
                refresh_sheet_list()
                sheet_listbox.select_set(idx)

        ctk.CTkButton(sheet_controls, text="Add Sheet", command=add_sheet, width=80, font=FONTS["button"]).pack(pady=2)
        ctk.CTkButton(sheet_controls, text="Rename", command=rename_sheet, width=80, font=FONTS["button"]).pack(pady=2)
        ctk.CTkButton(sheet_controls, text="Remove", command=remove_sheet, width=80,
                      fg_color="red", hover_color="#8b0000", font=FONTS["button"]).pack(pady=2)

        def add_col():
            idx = get_current_sheet_idx()
            if idx is None:
                return
            for i in listbox_avail.curselection():
                editing_config[idx]['columns'].append(listbox_avail.get(i))
            refresh_columns_view()

        def remove_col():
            idx = get_current_sheet_idx()
            if idx is None:
                return
            to_remove = {listbox_sel.get(i) for i in listbox_sel.curselection()}
            editing_config[idx]['columns'] = [c for c in editing_config[idx]['columns'] if c not in to_remove]
            refresh_columns_view()

        def add_all():
            idx = get_current_sheet_idx()
            if idx is None:
                return
            editing_config[idx]['columns'] = list(available_columns)
            refresh_columns_view()

        def remove_all():
            idx = get_current_sheet_idx()
            if idx is None:
                return
            editing_config[idx]['columns'] = []
            refresh_columns_view()

        col_buttons = []
        col_buttons.append(ctk.CTkButton(mid_frame, text="Add >", command=add_col, width=80, font=FONTS["button"]))
        col_buttons[-1].pack(pady=2)
        col_buttons.append(ctk.CTkButton(mid_frame, text="< Remove", command=remove_col, width=80, font=FONTS["button"]))
        col_buttons[-1].pack(pady=2)
        col_buttons.append(ctk.CTkButton(mid_frame, text="Add All >>", command=add_all, width=80, font=FONTS["button"]))
        col_buttons[-1].pack(pady=10)
        col_buttons.append(ctk.CTkButton(mid_frame, text="<< Clear", command=remove_all, width=80, font=FONTS["button"]))
        col_buttons[-1].pack(pady=2)

        def save():
            for sheet in editing_config:
                if any(s in sheet.get('criteria', []) for s in _EXCLUSIVE_STATS):
                    continue
                if not sheet['columns']:
                    messagebox.showwarning("Warning", f"Sheet '{sheet['name']}' has no columns selected.", parent=dialog)
                    return
                if not sheet.get('criteria') or all(c == 'Group By Criteria' for c in sheet.get('criteria', [])):
                    if not messagebox.askyesno("Warning",
                                               f"Sheet '{sheet['name']}' has no failure monitoring criteria.\n"
                                               "All test cases will be shown.\n\nContinue?", parent=dialog):
                        return
            self.custom_report_sheets = editing_config
            self.generate_custom_report = True
            dialog.destroy()

        def cancel():
            self.generate_custom_report = False
            dialog.destroy()

        ctk.CTkButton(btn_frame, text="Cancel", command=cancel,
                      fg_color="gray", hover_color="#404040", font=FONTS["button"]).pack(side="left", padx=10, expand=True)
        ctk.CTkButton(btn_frame, text="Save & Run", command=save,
                      fg_color="green", hover_color="#006400", font=FONTS["button"]).pack(side="left", padx=10, expand=True)

        refresh_sheet_list()
        dialog.protocol("WM_DELETE_WINDOW", cancel)
        self.frame.wait_window(dialog)

    def regenerate_custom_report(self):
        """Generate a custom report from the existing allure-results folder (no re-run needed)."""
        allure_dir = self.last_allure_dir or self.allure_results_dir.get()

        if not os.path.isdir(allure_dir):
            messagebox.showwarning(
                "No Allure Results",
                f"Allure results folder not found:\n{allure_dir}\n\nPlease run tests first."
            )
            return

        if not self.custom_report_sheets:
            self.show_custom_report_dialog()
            if not self.generate_custom_report:
                return

        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            default_name = f"custom_report_{timestamp}.xlsx"
            report_path = filedialog.asksaveasfilename(
                title="Save Custom Report",
                defaultextension=".xlsx",
                initialfile=default_name,
                initialdir=os.path.dirname(self.last_custom_report_path) if self.last_custom_report_path else None,
                filetypes=[("Excel files", "*.xlsx")]
            )
            if not report_path:
                self.status_var.set("Report generation cancelled")
                return

            self.status_var.set("Generating custom report...")
            os.environ[EnvVar.SOAP_EXECUTION_MODE] = ExecutionMode.SYSTEM1_ONLY
            generate_custom_combined_report(
                allure_dir,
                report_path,
                self.custom_report_sheets,
                assert_error_text_enabled=False
            )
            self.last_custom_report_path = report_path
            self.last_allure_dir = allure_dir
            self.status_var.set(f"Custom report generated: {os.path.basename(report_path)}")
            messagebox.showinfo("Excel Generated", f"Custom report generated at:\n{report_path}")

        except Exception as e:
            messagebox.showerror("Excel Error", f"Failed to generate Custom Report: {e}")
            self.status_var.set("Report generation failed")

    def show_pytest_output(self, stdout, stderr):
        output_window = ctk.CTkToplevel(self.frame)
        output_window.title("Pytest Execution Output")
        output_window.geometry("800x600")

        tabview = ctk.CTkTabview(output_window)
        tabview.pack(fill="both", expand=True, padx=10, pady=10)
        
        tabview.add("Standard Output")
        tabview.add("Error Output")

        stdout_text = ctk.CTkTextbox(tabview.tab("Standard Output"))
        stdout_text.pack(fill="both", expand=True)
        stdout_text.insert("1.0", stdout)
        stdout_text.configure(state='disabled')

        stderr_text = ctk.CTkTextbox(tabview.tab("Error Output"))
        stderr_text.pack(fill="both", expand=True)
        stderr_text.insert("1.0", stderr)
        stderr_text.configure(state='disabled')
        
        ctk.CTkButton(output_window, text="Close", command=output_window.destroy).pack(pady=5)

    def _clear_allure_results(self, allure_dir):
        self.allure_report_generated = False
        try:
            if os.path.exists(allure_dir):
                shutil.rmtree(allure_dir)
            os.makedirs(allure_dir, exist_ok=True)
        except Exception as e:
            print(f"WARNING: Could not clear allure results: {e}")

    def load_sheet_data(self, sheet_name):
        file_path = self.file_path_var.get()
        try:
            self.status_var.set(f"Loading sheet '{sheet_name}'...")
            self.frame.update()

            if sheet_name in self.all_sheets_data:
                self.all_data = self.all_sheets_data[sheet_name]
            else:
                self.all_data = self.excel_handler.load_excel(file_path, sheet_name=sheet_name)
                self.all_sheets_data[sheet_name] = self.all_data

            if self.all_data:
                total_rows = len(self.all_data)
                rows_per_page_str = self.rows_per_page_var.get()

                if rows_per_page_str == "All":
                    self.rows_per_page = total_rows
                else:
                    self.rows_per_page = int(rows_per_page_str)

                self.total_pages = (total_rows + self.rows_per_page - 1) // self.rows_per_page
                self.current_page = 0

                if total_rows > 100:
                    self.pagination_frame.pack(fill="x", pady=(0, 10), before=self.tree.master)
                else:
                    self.pagination_frame.pack_forget()

                self.display_current_page()
                status_msg = f"Loaded {total_rows} rows"
                self.status_var.set(status_msg)
            else:
                self.status_var.set("Loaded 0 rows")
                self.pagination_frame.pack_forget()

        except Exception as e:
            messagebox.showerror("Error", f"Failed to load sheet '{sheet_name}': {str(e)}")

    def display_current_page(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        if not self.all_data:
            return

        start_idx = self.current_page * self.rows_per_page
        end_idx = min(start_idx + self.rows_per_page, len(self.all_data))
        page_data = self.all_data[start_idx:end_idx]

        if not self.tree['columns']:
            columns = list(self.all_data[0].keys())
            self.tree['columns'] = columns
            for col in columns:
                self.tree.heading(col, text=col)
                self.tree.column(col, width=150)

        for row in page_data:
            columns = self.tree['columns']
            values = [row.get(col, '') for col in columns]
            self.tree.insert('', 'end', values=values)

        self.page_label.configure(text=f"Page {self.current_page + 1} of {self.total_pages} (Rows {start_idx + 1}-{end_idx} of {len(self.all_data)})")

    def on_rows_per_page_changed(self, event=None):
        rows_per_page_str = self.rows_per_page_var.get()
        if rows_per_page_str == "All":
            self.rows_per_page = len(self.all_data)
            self.pagination_frame.pack_forget()
        else:
            self.rows_per_page = int(rows_per_page_str)
            if len(self.all_data) > 100:
                self.pagination_frame.pack(fill="x", pady=(0, 10), before=self.tree.master)

        self.total_pages = max(1, (len(self.all_data) + self.rows_per_page - 1) // self.rows_per_page)
        self.current_page = 0
        self.display_current_page()

    def go_to_first_page(self):
        if self.current_page != 0:
            self.current_page = 0
            self.display_current_page()

    def go_to_prev_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.display_current_page()

    def go_to_next_page(self):
        if self.current_page < self.total_pages - 1:
            self.current_page += 1
            self.display_current_page()

    def go_to_last_page(self):
        if self.current_page != self.total_pages - 1:
            self.current_page = self.total_pages - 1
            self.display_current_page()

    def go_to_page(self, event=None):
        try:
            page_num = int(self.page_entry.get())
            if 1 <= page_num <= self.total_pages:
                self.current_page = page_num - 1
                self.display_current_page()
                self.page_entry.delete(0, tk.END)
            else:
                messagebox.showwarning("Invalid Page", f"Please enter a page between 1 and {self.total_pages}")
        except ValueError:
            messagebox.showwarning("Invalid Input", "Please enter a valid page number")

    def create_template(self):
        file_path = filedialog.asksaveasfilename(title="Save Template", defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if file_path:
            try:
                create_excel_template(file_path)
                messagebox.showinfo("Success", f"Template created at {file_path}")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to create template: {str(e)}")

    def save_allure_report(self):
        report_dir = "allure-report"
        results_dir = self.allure_results_dir.get()
        if not os.path.exists(results_dir):
            messagebox.showwarning("Error", "No test results found. Please run tests first.")
            return

        # Ask user for format preference
        save_single = messagebox.askyesno("Save Format", "Do you want to combine the report into a single HTML file?\n\nYes: Single .html file (easier to share)\nNo: Full folder (standard Allure format)")

        if save_single:
            # Allure natively supports single-file report generation
            # (allure generate --single-file), so no external dependency needed.
            allure_cmd = self._find_allure_executable()
            if not allure_cmd:
                messagebox.showerror("Error", "Allure not found.")
                return
            try:
                self.status_var.set("Generating single-file report...")
                self.frame.update()
                subprocess.run(
                    [allure_cmd, "generate", "--single-file", results_dir, "-o", report_dir, "--clean"],
                    capture_output=True, shell=False, check=True,
                )
                # This report_dir now holds the single-file variant, not the
                # regular multi-file report, so it must be regenerated next
                # time it's needed as a folder.
                self.allure_report_generated = False

                source_file = os.path.join(report_dir, "index.html")
                if os.path.exists(source_file):
                    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                    default_name = f"allure-report_{timestamp}.html"

                    file_path = filedialog.asksaveasfilename(
                        title="Save Single HTML Report",
                        defaultextension=".html",
                        initialfile=default_name,
                        filetypes=[("HTML files", "*.html")]
                    )

                    if file_path:
                        shutil.copy2(source_file, file_path)
                        messagebox.showinfo("Success", f"Report saved to:\n{file_path}")
                        self.status_var.set("Report saved")
                    else:
                        self.status_var.set("Save cancelled")
                else:
                    messagebox.showerror("Error", "Failed to generate single file (index.html NOT found).")
                    self.status_var.set("Generation failed")

            except Exception as e:
                messagebox.showerror("Error", f"Failed to generate single-file report: {e}")
                self.status_var.set("Error generating report")

        else:
            # Regenerate HTML only if tests were run since the last generation
            if not self.allure_report_generated:
                allure_cmd = self._find_allure_executable()
                if not allure_cmd:
                    messagebox.showerror("Error", "Allure not found.")
                    return
                self.status_var.set("Generating HTML report from latest results...")
                self.frame.update()
                subprocess.run([allure_cmd, "generate", results_dir, "-o", report_dir, "--clean"], capture_output=True, shell=False)
                self.allure_report_generated = True

            destination = filedialog.askdirectory(title="Select Directory")
            if destination:
                try:
                    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                    folder_name = f"allure-report_{timestamp}"
                    dest_path = os.path.join(destination, folder_name)
                    shutil.copytree(report_dir, dest_path)
                    if messagebox.askyesno("Zip Report", "Do you want to zip the report folder?"):
                        self.status_var.set("Zipping report...")
                        self.frame.update()
                        zip_path = shutil.make_archive(dest_path, 'zip', destination, folder_name)
                        shutil.rmtree(dest_path)
                        messagebox.showinfo("Success", f"Report saved and zipped to:\n{zip_path}")
                        self.status_var.set("Report saved and zipped")
                    else:
                        messagebox.showinfo("Success", f"Report saved to:\n{dest_path}")
                        self.status_var.set("Report saved")
                except Exception as e:
                    messagebox.showerror("Error", f"Failed to save: {e}")

    def open_allure_report(self):
        results_dir = os.path.abspath(self.allure_results_dir.get())
        if not os.path.exists(results_dir):
            messagebox.showwarning("Error", "No test results found.")
            return

        allure_cmd = self._find_allure_executable()
        if allure_cmd:
            subprocess.Popen([allure_cmd, "serve", results_dir], shell=False, cwd=tempfile.gettempdir())
        else:
            messagebox.showerror("Error", "Allure not found")

    def generate_allure_html_report_manual(self):
        self.generate_allure_html_report(self.allure_results_dir.get())

    def generate_allure_html_report(self, results_dir):
        allure_cmd = self._find_allure_executable()
        if not allure_cmd:
             messagebox.showerror("Error", "Allure not found")
             return
        
        report_dir = "allure-report"
        subprocess.run([allure_cmd, "generate", results_dir, "-o", report_dir, "--clean"], capture_output=True, shell=False)
        self.status_var.set(f"HTML report generated in {report_dir}")
        messagebox.showinfo("Success", "HTML generated")

    def _find_allure_executable(self):
        import shutil
        path = shutil.which("allure")
        if path: return path
        for name in ("allure.bat", "allure.cmd"):
            path = shutil.which(name)
            if path: return path
        return None

    def stop_tests(self):
        self.testing_active = False
        self.status_var.set("Stopping...")
        if self.current_process:
            self.current_process.terminate()