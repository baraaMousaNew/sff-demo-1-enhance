"""
Request Generator Tab - Load an Excel sheet, map columns, execute requests, and write responses back.
"""

import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import customtkinter as ctk
import openpyxl

from utils.api_client import APIClient
from utils.excel_handler import ExcelHandler
from utils.request_sender.soap_action_enums import build_soap_action
from .theme import FONTS


class ToolTip:
    def __init__(self, widget, text):
        self.widget = widget
        self.text = text
        self.tooltip_window = None
        self.widget.bind("<Enter>", self.show_tooltip)
        self.widget.bind("<Leave>", self.hide_tooltip)

    def show_tooltip(self, event=None):
        if self.tooltip_window:
            return
        x = self.widget.winfo_rootx() + 20
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 10
        self.tooltip_window = tk.Toplevel(self.widget)
        self.tooltip_window.wm_overrideredirect(True)
        frame = tk.Frame(self.tooltip_window, background="#ffffe0", relief="solid", borderwidth=1)
        frame.pack()
        tk.Label(frame, text=self.text, background="#ffffe0", foreground="black",
                 justify="left", font=("Arial", 10)).pack(padx=5, pady=2)
        self.tooltip_window.wm_geometry(f"+{x}+{y}")

    def hide_tooltip(self, event=None):
        if self.tooltip_window:
            self.tooltip_window.destroy()
            self.tooltip_window = None


class RequestGeneratorTab:
    def __init__(self, parent):
        self.frame = ctk.CTkFrame(parent, fg_color="transparent")
        self.frame.pack(fill="both", expand=True)

        self.excel_handler = ExcelHandler()
        self.api_client = APIClient()

        self.sheet_names = []
        self.current_sheet = None
        self.all_data = []          # list of dicts from pandas
        self.columns = []           # column names of the loaded sheet
        self.executing = False
        self.last_summary = []      # kept so the summary dialog can be reopened

        self.create_widgets()

    # ------------------------------------------------------------------
    # Widget construction
    # ------------------------------------------------------------------

    def create_widgets(self):
        main_frame = ctk.CTkScrollableFrame(self.frame)
        main_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # ── 1. File selection ──────────────────────────────────────────
        self.file_frame = ctk.CTkFrame(main_frame)
        file_frame = self.file_frame
        file_frame.pack(fill="x", pady=(0, 8))

        ctk.CTkLabel(file_frame, text="Excel File:", font=FONTS["main_bold"]).pack(side="left", padx=10)
        self.file_path_var = tk.StringVar()
        ctk.CTkEntry(file_frame, textvariable=self.file_path_var, width=400, font=FONTS["main"]).pack(
            side="left", fill="x", expand=True, padx=5
        )
        ctk.CTkButton(
            file_frame, text="Browse", command=self.browse_file,
            width=80, fg_color="gray", hover_color="#404040", font=FONTS["button"]
        ).pack(side="left", padx=(0, 5))
        ctk.CTkButton(
            file_frame, text="Load", command=self.load_file,
            width=80, font=FONTS["button"]
        ).pack(side="left", padx=5)

        # ── 2. Sheet selector (hidden until multi-sheet file) ──────────
        self.sheet_frame = ctk.CTkFrame(main_frame)
        ctk.CTkLabel(self.sheet_frame, text="Sheet:", font=FONTS["main_bold"]).pack(side="left", padx=10)
        self.sheet_var = tk.StringVar()
        self.sheet_combo = ctk.CTkComboBox(
            self.sheet_frame, variable=self.sheet_var, values=[],
            width=200, font=FONTS["main"], command=self.on_sheet_changed,
        )
        self.sheet_combo.pack(side="left", padx=5)

        # ── 3. Endpoint configuration ──────────────────────────────────
        endpoint_frame = ctk.CTkFrame(main_frame)
        endpoint_frame.pack(fill="x", pady=(0, 8))

        ctk.CTkLabel(endpoint_frame, text="URL:", font=FONTS["main_bold"]).pack(side="left", padx=10)
        self.url_var = tk.StringVar()
        ctk.CTkEntry(
            endpoint_frame, textvariable=self.url_var, width=500,
            placeholder_text="https://...", font=FONTS["main"]
        ).pack(side="left", fill="x", expand=True, padx=5)

        # ── 4. Column mapping (hidden until a sheet is loaded) ─────────
        self.mapping_outer = ctk.CTkFrame(main_frame)

        ctk.CTkLabel(
            self.mapping_outer, text="Column Mapping", font=FONTS["sub_header"]
        ).pack(anchor="w", padx=10, pady=(8, 4))

        mapping_inner = ctk.CTkFrame(self.mapping_outer, fg_color="transparent")
        mapping_inner.pack(fill="x", padx=10, pady=(0, 8))

        def _mapping_row(label, row):
            ctk.CTkLabel(mapping_inner, text=label, font=FONTS["main_bold"], width=160, anchor="w").grid(
                row=row, column=0, padx=(0, 10), pady=4, sticky="w"
            )
            var = tk.StringVar()
            combo = ctk.CTkComboBox(mapping_inner, variable=var, values=[], width=220, font=FONTS["main"])
            combo.grid(row=row, column=1, padx=(0, 20), pady=4, sticky="w")
            return var, combo

        self.transaction_col_var, self.transaction_col_combo = _mapping_row("Transaction Name:", 0)
        self.payload_col_var,     self.payload_col_combo     = _mapping_row("Request / Payload:", 1)
        self.response_col_var,    self.response_col_combo    = _mapping_row("Response (write to):", 2)

        # ── 5. Preview table ───────────────────────────────────────────
        ctk.CTkLabel(main_frame, text="File Preview:", font=FONTS["sub_header"]).pack(anchor="w", padx=5, pady=(4, 2))

        tree_frame = ctk.CTkFrame(main_frame)
        tree_frame.pack(fill="both", expand=True, pady=(0, 8))

        self.tree = ttk.Treeview(tree_frame, show="headings", height=14, selectmode="extended")
        v_scroll = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL,   command=self.tree.yview)
        h_scroll = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)
        self.tree.pack(side="left", fill="both", expand=True)
        v_scroll.pack(side="right",  fill="y")
        h_scroll.pack(side="bottom", fill="x")
        self.tree.bind("<Double-1>", self._on_cell_double_click)

        # ── 6. Progress & status ───────────────────────────────────────
        progress_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        progress_frame.pack(fill="x", pady=(0, 4))

        self.progress_var = tk.DoubleVar(value=0)
        self.progress_bar = ctk.CTkProgressBar(progress_frame, variable=self.progress_var)
        self.progress_bar.pack(fill="x", pady=(4, 2))
        self.progress_bar.set(0)

        self.status_var = tk.StringVar(value="Ready")
        ctk.CTkLabel(progress_frame, textvariable=self.status_var, text_color="cyan", font=FONTS["main"]).pack(
            anchor="w", pady=(0, 4)
        )

        # ── 7. Action buttons ──────────────────────────────────────────
        btn_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        btn_frame.pack(fill="x", pady=6)

        _execution_tooltip = "Rows with no value in the Request/Payload column will be skipped."

        self.execute_btn = ctk.CTkButton(
            btn_frame, text="Execute All", command=self.start_execution,
            fg_color="green", hover_color="#006400", width=140, font=FONTS["button"]
        )
        self.execute_btn.pack(side="left", padx=(0, 10))
        ToolTip(self.execute_btn, _execution_tooltip)

        self.execute_selected_btn = ctk.CTkButton(
            btn_frame, text="Execute Selected", command=self.start_execution_selected,
            fg_color="#1f7a1f", hover_color="#145214", width=160, font=FONTS["button"]
        )
        self.execute_selected_btn.pack(side="left", padx=(0, 10))
        ToolTip(self.execute_selected_btn, _execution_tooltip)

        self.stop_btn = ctk.CTkButton(
            btn_frame, text="Stop", command=self.stop_execution,
            fg_color="red", hover_color="#8b0000", width=100, font=FONTS["button"],
            state="disabled"
        )
        self.stop_btn.pack(side="left", padx=(0, 10))

        ctk.CTkButton(
            btn_frame, text="Save Responses to Excel", command=self.save_responses,
            fg_color="#1f538d", font=FONTS["button"]
        ).pack(side="left", padx=(0, 10))

        self.summary_btn = ctk.CTkButton(
            btn_frame, text="View Last Summary", command=self.open_last_summary,
            fg_color="gray", hover_color="#404040", font=FONTS["button"], state="disabled"
        )
        self.summary_btn.pack(side="left")

    # ------------------------------------------------------------------
    # File handling
    # ------------------------------------------------------------------

    def browse_file(self):
        path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls")],
        )
        if path:
            self.file_path_var.set(path)
            self.detect_sheets()

    def detect_sheets(self):
        path = self.file_path_var.get()
        if not path:
            return
        try:
            self.sheet_names = self.excel_handler.get_sheet_names(path)
            if len(self.sheet_names) > 1:
                self.sheet_combo.configure(values=self.sheet_names)
                self.sheet_combo.set(self.sheet_names[0])
                self.current_sheet = self.sheet_names[0]
                self.sheet_frame.pack(fill="x", pady=(0, 8), after=self.file_frame)
                self.status_var.set(f"{len(self.sheet_names)} sheets found — select one and click Load.")
            else:
                self.sheet_frame.pack_forget()
                self.current_sheet = self.sheet_names[0] if self.sheet_names else None
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read sheets: {e}")

    def load_file(self):
        path = self.file_path_var.get()
        if not path:
            messagebox.showwarning("Warning", "Please select a file first.")
            return
        if not self.sheet_names:
            self.detect_sheets()
        sheet = self.current_sheet or (self.sheet_names[0] if self.sheet_names else None)
        if not sheet:
            messagebox.showerror("Error", "No sheets found in the Excel file.")
            return
        self.load_sheet_data(sheet)

    def on_sheet_changed(self, value):
        self.current_sheet = value
        self.load_sheet_data(value)

    def load_sheet_data(self, sheet_name):
        path = self.file_path_var.get()
        try:
            self.status_var.set(f"Loading '{sheet_name}'…")
            self.frame.update()
            self.all_data = self.excel_handler.load_excel(path, sheet_name=sheet_name)
            self.columns = list(self.all_data[0].keys()) if self.all_data else []
            self._update_column_mapping()
            self.populate_table()
            self.status_var.set(f"Loaded {len(self.all_data)} rows from '{sheet_name}'.")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load '{sheet_name}': {e}")

    # ------------------------------------------------------------------
    # Column mapping
    # ------------------------------------------------------------------

    def _update_column_mapping(self):
        """Populate the mapping dropdowns with the current sheet's column names."""
        for combo, var in [
            (self.transaction_col_combo, self.transaction_col_var),
            (self.payload_col_combo,     self.payload_col_var),
            (self.response_col_combo,    self.response_col_var),
        ]:
            combo.configure(values=self.columns)
            if self.columns and not var.get():
                var.set(self.columns[0])

        self.mapping_outer.pack(fill="x", pady=(0, 8))

    # ------------------------------------------------------------------
    # Table rendering
    # ------------------------------------------------------------------

    def populate_table(self):
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = []

        if not self.all_data:
            return

        display_columns = ["#"] + self.columns
        self.tree["columns"] = display_columns
        self.tree.heading("#", text="#")
        self.tree.column("#", width=50, minwidth=40, stretch=False, anchor="center")
        for col in self.columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=150, minwidth=80)

        for i, row in enumerate(self.all_data, start=1):
            self.tree.insert("", "end", values=[i] + [row.get(col, "") for col in self.columns])

    def _refresh_tree_row(self, row_index, row_data):
        """Update a single Treeview row after a response is written."""
        items = self.tree.get_children()
        if row_index < len(items):
            self.tree.item(items[row_index], values=[row_index + 1] + [row_data.get(col, "") for col in self.columns])

    def _on_cell_double_click(self, event):
        row_id = self.tree.identify_row(event.y)
        col_id = self.tree.identify_column(event.x)
        if not row_id or not col_id:
            return
        col_index = int(col_id.replace("#", "")) - 1
        # col_index 0 is the row-number column — nothing to expand
        if col_index <= 0 or col_index > len(self.columns):
            return
        col_name = self.columns[col_index - 1]
        value = self.tree.item(row_id, "values")[col_index]
        self._show_cell_dialog(col_name, value)

    def _show_cell_dialog(self, col_name, value):
        dialog = ctk.CTkToplevel(self.frame)
        dialog.title(col_name)
        dialog.geometry("800x500")
        dialog.transient(self.frame)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text=col_name, font=FONTS["sub_header"]).pack(anchor="w", padx=12, pady=(10, 4))

        text = ctk.CTkTextbox(dialog, font=FONTS["code"], wrap="none")
        text.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        text.insert("1.0", value)
        text.configure(state="disabled")

        ctk.CTkButton(
            dialog, text="Close", command=dialog.destroy,
            fg_color="gray", hover_color="#404040", font=FONTS["button"], width=100
        ).pack(pady=(0, 10))

    # ------------------------------------------------------------------
    # Execution
    # ------------------------------------------------------------------

    def _validate_execution_preconditions(self):
        """Returns True and clears to proceed, or shows a warning and returns False."""
        if not self.all_data:
            messagebox.showwarning("Warning", "Please load a file first.")
            return False
        if not self.url_var.get().strip():
            messagebox.showwarning("Warning", "Please enter a URL.")
            return False
        if not self.payload_col_var.get() or not self.response_col_var.get():
            messagebox.showwarning("Warning", "Please select the Payload and Response columns.")
            return False
        if self.response_col_var.get() not in self.columns:
            messagebox.showwarning("Warning", f"Response column '{self.response_col_var.get()}' not found in sheet.")
            return False
        return True

    def _begin_execution(self, rows):
        """Disable controls and launch the background thread for the given (index, row) pairs."""
        self.executing = True
        self.execute_btn.configure(state="disabled")
        self.execute_selected_btn.configure(state="disabled")
        self.stop_btn.configure(state="normal")
        self.progress_bar.set(0)
        thread = threading.Thread(target=self._run_execution, args=(rows,), daemon=True)
        thread.start()

    def start_execution(self):
        if not self._validate_execution_preconditions():
            return
        if not messagebox.askyesno("Confirm", f"Execute requests for all {len(self.all_data)} rows?"):
            return
        self._begin_execution(list(enumerate(self.all_data)))

    def start_execution_selected(self):
        if not self._validate_execution_preconditions():
            return
        all_items = self.tree.get_children()
        selected_items = self.tree.selection()
        if not selected_items:
            messagebox.showwarning("Warning", "No rows selected. Click rows in the table to select them.")
            return
        indices = [all_items.index(item) for item in selected_items]
        rows = [(i, self.all_data[i]) for i in indices]
        if not messagebox.askyesno("Confirm", f"Execute requests for {len(rows)} selected row(s)?"):
            return
        self._begin_execution(rows)

    def stop_execution(self):
        self.executing = False
        self.status_var.set("Stopping…")

    def _run_execution(self, rows):
        url             = self.url_var.get().strip()
        payload_col     = self.payload_col_var.get()
        transaction_col = self.transaction_col_var.get()
        response_col    = self.response_col_var.get()
        total           = len(rows)
        summary         = []   # list of dicts, one per row

        for progress, (i, row) in enumerate(rows):
            if not self.executing:
                summary.append({
                    "row": i + 1,
                    "transaction": str(row.get(transaction_col, "")),
                    "soap_action": "",
                    "status": "Aborted",
                    "detail": "Execution stopped by user",
                })
                for _, (j, r) in enumerate(rows[progress + 1:]):
                    summary.append({
                        "row": j + 1,
                        "transaction": str(r.get(transaction_col, "")),
                        "soap_action": "",
                        "status": "Aborted",
                        "detail": "Execution stopped by user",
                    })
                break

            transaction_name = str(row.get(transaction_col, "")).strip()

            if str(row.get(response_col, "")).strip():
                summary.append({
                    "row": i + 1,
                    "transaction": transaction_name,
                    "soap_action": "",
                    "status": "Skipped",
                    "detail": "Response column already filled",
                })
                self.frame.after(0, self.progress_var.set, (progress + 1) / total)
                self.frame.after(0, self.status_var.set,
                    f"Row {progress + 1} / {total} — skipped (response already filled)")
                continue

            payload = str(row.get(payload_col, "")).strip()
            if not payload:
                summary.append({
                    "row": i + 1,
                    "transaction": transaction_name,
                    "soap_action": "",
                    "status": "Skipped",
                    "detail": "Request/Payload column is empty",
                })
                self.frame.after(0, self.progress_var.set, (progress + 1) / total)
                self.frame.after(0, self.status_var.set,
                    f"Row {progress + 1} / {total} — skipped (no payload)")
                continue

            soap_action = build_soap_action(transaction_name, payload)
            if soap_action is None:
                summary.append({
                    "row": i + 1,
                    "transaction": transaction_name,
                    "soap_action": "",
                    "status": "Skipped",
                    "detail": f"Unrecognised transaction name: '{transaction_name}'",
                })
                self.frame.after(0, self.progress_var.set, (progress + 1) / total)
                self.frame.after(0, self.status_var.set,
                    f"Row {progress + 1} / {total} — skipped (unrecognised transaction)")
                continue

            try:
                result = self.api_client.send_request(
                    url=url,
                    xml_body=payload,
                    soap_action=soap_action,
                )
                row[response_col] = result["content"]
                summary.append({
                    "row": i + 1,
                    "transaction": transaction_name,
                    "soap_action": soap_action,
                    "status": "Success" if result["success"] else "HTTP Error",
                    "detail": f"HTTP {result['status_code']} — {result['response_time']:.0f} ms",
                })
            except Exception as e:
                row[response_col] = f"ERROR: {e}"
                summary.append({
                    "row": i + 1,
                    "transaction": transaction_name,
                    "soap_action": soap_action,
                    "status": "Error",
                    "detail": str(e),
                })

            self.frame.after(0, self._refresh_tree_row, i, row)
            self.frame.after(0, self.progress_var.set, (progress + 1) / total)
            self.frame.after(0, self.status_var.set,
                f"Row {progress + 1} / {total} — {transaction_name}")

        self.frame.after(0, self._on_execution_done, summary)

    def _on_execution_done(self, summary):
        self.executing = False
        self.last_summary = summary
        self.execute_btn.configure(state="normal")
        self.execute_selected_btn.configure(state="normal")
        self.stop_btn.configure(state="disabled")
        self.summary_btn.configure(state="normal")
        self.status_var.set("Done. Click 'Save Responses to Excel' to write results back to the file.")
        self._show_summary_dialog(summary)

    def open_last_summary(self):
        if self.last_summary:
            self._show_summary_dialog(self.last_summary)

    def _show_summary_dialog(self, summary):
        dialog = ctk.CTkToplevel(self.frame)
        dialog.title("Execution Summary")
        dialog.geometry("950x450")
        dialog.transient(self.frame)
        dialog.grab_set()

        # counts
        counts = {"Success": 0, "HTTP Error": 0, "Error": 0, "Skipped": 0, "Aborted": 0}
        for r in summary:
            counts[r["status"]] = counts.get(r["status"], 0) + 1

        header = (
            f"Total: {len(summary)}    "
            f"Success: {counts['Success']}    "
            f"Skipped: {counts['Skipped']}    "
            f"Error: {counts['Error'] + counts['HTTP Error']}    "
            f"Aborted: {counts['Aborted']}"
        )
        ctk.CTkLabel(dialog, text=header, font=FONTS["main_bold"]).pack(anchor="w", padx=12, pady=(10, 6))

        # table
        table_frame = ctk.CTkFrame(dialog)
        table_frame.pack(fill="both", expand=True, padx=12, pady=(0, 8))

        columns = ("Row", "Transaction", "SOAP Action", "Status", "Detail")
        tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=14)
        tree.heading("Row",         text="Row",         anchor="center")
        tree.heading("Transaction", text="Transaction", anchor="w")
        tree.heading("SOAP Action", text="SOAP Action", anchor="w")
        tree.heading("Status",      text="Status",      anchor="center")
        tree.heading("Detail",      text="Detail",      anchor="w")
        tree.column("Row",         width=55,  anchor="center", stretch=False)
        tree.column("Transaction", width=160, anchor="w")
        tree.column("SOAP Action", width=260, anchor="w")
        tree.column("Status",      width=100, anchor="center", stretch=False)
        tree.column("Detail",      width=200, anchor="w")

        v_scroll = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=v_scroll.set)
        tree.pack(side="left", fill="both", expand=True)
        v_scroll.pack(side="right", fill="y")

        # colour tags
        tree.tag_configure("Success",    foreground="#2cc985")
        tree.tag_configure("HTTP Error", foreground="#ffc107")
        tree.tag_configure("Error",      foreground="#ff4d4d")
        tree.tag_configure("Skipped",    foreground="#888888")
        tree.tag_configure("Aborted",    foreground="#ff4d4d")

        for r in summary:
            tree.insert("", "end",
                values=(r["row"], r["transaction"], r["soap_action"], r["status"], r["detail"]),
                tags=(r["status"],))

        ctk.CTkButton(
            dialog, text="Close", command=dialog.destroy,
            fg_color="gray", hover_color="#404040", font=FONTS["button"], width=100
        ).pack(pady=(0, 10))

    # ------------------------------------------------------------------
    # Save responses back to Excel
    # ------------------------------------------------------------------

    def save_responses(self):
        if not self.all_data:
            messagebox.showwarning("Warning", "No data to save.")
            return

        src_path = self.file_path_var.get()
        response_col = self.response_col_var.get()

        if not response_col:
            messagebox.showwarning("Warning", "Please select the Response column first.")
            return

        save_path = filedialog.asksaveasfilename(
            title="Save Excel with Responses",
            initialfile=src_path,
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not save_path:
            return

        try:
            wb = openpyxl.load_workbook(src_path)
            sheet_name = self.current_sheet or wb.sheetnames[0]
            ws = wb[sheet_name]

            # Find header row and the response column index
            headers = [cell.value for cell in ws[1]]
            if response_col not in headers:
                messagebox.showerror("Error", f"Column '{response_col}' not found in the worksheet headers.")
                return
            col_idx = headers.index(response_col) + 1  # openpyxl is 1-based

            # Write responses (data starts at row 2)
            for row_num, row_data in enumerate(self.all_data, start=2):
                ws.cell(row=row_num, column=col_idx, value=row_data.get(response_col, ""))

            wb.save(save_path)
            messagebox.showinfo("Saved", f"Responses saved to:\n{save_path}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save file: {e}")
