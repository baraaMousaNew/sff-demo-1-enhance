"""
Excel Handler - Enhanced with multi-sheet support and Response Extraction column support
"""

import pandas as pd
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from utils import common_variables
from utils.common_variables import tc_precondition


def validate_excel(sheets_to_test: dict):
    expected_columns = [common_variables.tc_id, common_variables.tc_description, common_variables.tc_test_data,
                        common_variables.tc_rule_id, common_variables.tc_name, common_variables.tc_request_name, common_variables.tc_precondition,
                        common_variables.tc_objective, common_variables.tc_scenario_type, common_variables.tc_error_text]

    # messages = []
    for sheet in sheets_to_test:
        actual_columns = list(sheets_to_test[sheet][0].keys())
        missing_columns = set(expected_columns) - set(actual_columns)
        # extra_columns = set(actual_columns) - set(expected_columns)
        if missing_columns:
            return False, f"Sheet {sheet} has missing columns: {missing_columns}"
        # elif extra_columns:
        #     messages.append(f"Sheet {sheet} has extra columns: {extra_columns}")
    return True, None

def validate_assertion_excel(sheets_to_test: dict):
    expected_columns = [common_variables.tc_id, common_variables.tc_description, common_variables.tc_test_data,
                        common_variables.tc_rule_id, common_variables.tc_name, common_variables.tc_request_name, common_variables.tc_precondition,
                        common_variables.tc_objective, common_variables.tc_scenario_type]

    # messages = []
    for sheet in sheets_to_test:
        actual_columns = list(sheets_to_test[sheet][0].keys())
        missing_columns = set(expected_columns) - set(actual_columns)
        # extra_columns = set(actual_columns) - set(expected_columns)
        if missing_columns:
            return False, f"Sheet {sheet} has missing columns: {missing_columns}"
        # elif extra_columns:
        #     messages.append(f"Sheet {sheet} has extra columns: {extra_columns}")
    return True, None


def create_excel_template(file_path):
    """
    Create an Excel template with just the headers from the original file
    """

    # Define the headers as they appear in the original file
    headers = [
        'TC ID',
        'Transaction Type',
        'Rule ID',
        'Occurrences',
        'Precondition',
        'Object',
        'Element',
        'Title ',  # Note: original has a space after Title
        'Name',
        'Description',
        'Test Data',
        'Actual Result',
        'Objective (Expected Result)',
        'Priority',
        'ScenarioType',
        'Notes'
    ]

    # Create a new workbook and select the active worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"

    # Add headers to the first row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)

        # Format headers with bold font and light blue background
        cell.font = Font(bold=True, color="000000")
        cell.fill = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths based on header length
    for col_num, header in enumerate(headers, 1):
        column_letter = ws.cell(row=1, column=col_num).column_letter
        # Set minimum width of 12, maximum of 25
        width = max(12, min(len(header) + 2, 25))
        ws.column_dimensions[column_letter].width = width

    # Save the workbook
    wb.save(file_path)
    print(f"Excel template created successfully: {file_path}")
    print(f"Template contains {len(headers)} columns with headers only")

    return file_path

class ExcelHandler:
    def __init__(self):
        pass

    def get_sheet_names(self, file_path):
        """
        Get all sheet names from an Excel file

        Args:
            file_path (str): Path to Excel file

        Returns:
            list: List of sheet names
        """
        try:
            excel_file = pd.ExcelFile(file_path)
            return excel_file.sheet_names
        except Exception as e:
            raise Exception(f"Failed to read Excel file sheets: {str(e)}")

    def load_excel(self, file_path, sheet_name=None):
        """
        Load Excel file and return data as list of dictionaries

        Args:
            file_path (str): Path to Excel file
            sheet_name (str, optional): Name of sheet to load. If None, loads first sheet.

        Returns:
            list: List of row dictionaries
        """
        try:
            if sheet_name:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
            else:
                df = pd.read_excel(file_path)
            # Convert NaN to empty strings
            df = df.fillna('')
            return df.to_dict('records')

        except Exception as e:
            raise Exception(f"Failed to load Excel file: {str(e)}")


